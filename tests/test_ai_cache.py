"""Phase 12 tests: fingerprint cache, call budgets, logging and reports. FAKE client only."""

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

import config.settings as config_settings
import database.connection as dbc
from ai.cache import budget_status, compute_fingerprint, get_insights, today_start_utc
from ai.client import classify_error
from config.settings import Settings
from dashboard.pipeline import SAMPLE_DIR, load_raw, prepare, run_pipeline
from database import repository as repo
from reports.excel_report import generate_workbook
from reports.pdf_report import AI_NOT_GENERATED, generate_pdf
from tests.test_ai import FakeClient, good_answer, quota_error

ROOT = Path(__file__).resolve().parent.parent


def ai_settings(**overrides):
    base = dict(ai_enabled=True, gemini_api_key="k", gemini_model="fake-model",
                ai_max_calls_per_run=3, ai_max_calls_per_day=15, ai_cache_enabled=True)
    base.update(overrides)
    return Settings(**base)


@pytest.fixture(scope="module")
def outputs(tmp_path_factory):
    db = tmp_path_factory.mktemp("prep") / "p.db"
    out = {}
    for name in ("marketing_clean.csv", "test_c_large.csv"):
        raw, report = load_raw(SAMPLE_DIR / name)
        out[name] = run_pipeline(prepare(raw, report), db_path=db)
    return out


@pytest.fixture
def clean(outputs):
    return outputs["marketing_clean.csv"].analysis


class CountingFactory:
    """Builds FakeClients and counts how often a real client was needed."""

    def __init__(self, responses):
        self.client = FakeClient(responses)
        self.built = 0

    def __call__(self):
        self.built += 1
        return self.client


def answer_for(result) -> str:
    from ai.context_builder import build_evidence
    return json.dumps(good_answer(build_evidence(result, model="fake-model")))


# --- Fingerprint ------------------------------------------------------------------------------

def test_fingerprint_is_canonical_and_sensitive():
    a = compute_fingerprint('{"b": 1, "a": [1, 2]}', "m")
    assert a == compute_fingerprint('{"a": [1, 2], "b": 1}', "m")          # key order ignored
    assert a != compute_fingerprint('{"a": [1, 2], "b": 2}', "m")          # content matters
    assert a != compute_fingerprint('{"a": [1, 2], "b": 1}', "other")      # model matters
    assert a != compute_fingerprint('{"a": [1, 2], "b": 1}', "m", prompt_version="v99")


# --- Cache ------------------------------------------------------------------------------------

def test_same_context_twice_makes_one_call(clean, tmp_path):
    db = tmp_path / "c.db"
    factory = CountingFactory([answer_for(clean)])
    first = get_insights(clean, ai_settings(), factory, db_path=db)
    second = get_insights(clean, ai_settings(), factory, db_path=db)
    assert first.ok and not first.from_cache
    assert second.ok and second.from_cache and second.cached_at
    assert len(factory.client.calls) == 1 and factory.built == 1
    assert second.insights == first.insights
    usage = repo.ai_usage_since(today_start_utc(), db_path=db)
    assert usage["calls"] == 1 and usage["cache_hits"] == 1


def test_changed_data_means_new_fingerprint_and_call(outputs, tmp_path):
    db = tmp_path / "d.db"
    a, b = outputs["marketing_clean.csv"].analysis, outputs["test_c_large.csv"].analysis
    fa = CountingFactory([answer_for(a)])
    fb = CountingFactory([answer_for(b)])
    ra = get_insights(a, ai_settings(), fa, db_path=db)
    rb = get_insights(b, ai_settings(), fb, db_path=db)
    assert ra.pack.fingerprint != rb.pack.fingerprint
    assert len(fa.client.calls) == 1 and len(fb.client.calls) == 1 and not rb.from_cache


def test_cache_can_be_disabled(clean, tmp_path):
    db = tmp_path / "n.db"
    factory = CountingFactory([answer_for(clean), answer_for(clean)])
    get_insights(clean, ai_settings(ai_cache_enabled=False), factory, db_path=db)
    again = get_insights(clean, ai_settings(ai_cache_enabled=False), factory, db_path=db)
    assert len(factory.client.calls) == 2 and not again.from_cache


def test_regenerate_skips_cache_and_uses_budget(clean, tmp_path):
    db = tmp_path / "f.db"
    factory = CountingFactory([answer_for(clean), answer_for(clean)])
    get_insights(clean, ai_settings(), factory, db_path=db)
    forced = get_insights(clean, ai_settings(), factory, force=True, db_path=db)
    assert len(factory.client.calls) == 2 and not forced.from_cache
    assert repo.count_ai_calls_since(today_start_utc(), db_path=db) == 2


def test_cache_only_lookup_never_builds_a_client(clean, tmp_path):
    factory = CountingFactory([])
    assert get_insights(clean, ai_settings(), factory, allow_call=False, db_path=tmp_path / "e.db") is None
    assert factory.built == 0
    # Cached insights are readable even when AI is switched off (e.g. for reports).
    db = tmp_path / "g.db"
    get_insights(clean, ai_settings(), CountingFactory([answer_for(clean)]), db_path=db)
    off = get_insights(clean, ai_settings(ai_enabled=False), factory, allow_call=False, db_path=db)
    assert off is not None and off.from_cache and factory.built == 0


# --- Budget -----------------------------------------------------------------------------------

def test_daily_budget_reached_means_no_call(clean, tmp_path):
    db = tmp_path / "b.db"
    settings = ai_settings(ai_max_calls_per_day=1)
    first = CountingFactory([answer_for(clean)])
    get_insights(clean, settings, first, db_path=db)
    blocked = CountingFactory([answer_for(clean)])
    run = get_insights(clean, settings, blocked, force=True, db_path=db)
    assert blocked.built == 0 and run.error.kind == "budget"
    assert "daily AI limit" in run.error.user_message
    # The saved answer is still served without a call.
    cached = get_insights(clean, settings, blocked, db_path=db)
    assert cached.ok and cached.from_cache


def test_per_run_budget(outputs, tmp_path):
    db = tmp_path / "r.db"
    out = outputs["test_c_large.csv"]
    ds = repo.register_dataset("x.csv", "hash-r", db_path=db)
    run_id = repo.create_run(ds.id, db_path=db)
    settings = ai_settings(ai_max_calls_per_run=1)
    get_insights(out.analysis, settings, CountingFactory([answer_for(out.analysis)]), run_id=run_id, db_path=db)
    status = budget_status(settings, run_id, db_path=db)
    assert status.calls_this_run == 1 and status.remaining == 0
    blocked = get_insights(out.analysis, settings, CountingFactory([]), run_id=run_id, force=True, db_path=db)
    assert blocked.error.kind == "budget" and "this analysis run" in blocked.error.user_message


def test_unsaved_run_uses_session_counter(clean, tmp_path):
    status = budget_status(ai_settings(ai_max_calls_per_run=2), None, session_calls=2, db_path=tmp_path / "s.db")
    assert status.remaining == 0 and status.reason


def test_invalid_json_retry_capped_by_budget(clean, tmp_path):
    factory = CountingFactory(["not json", answer_for(clean)])
    run = get_insights(clean, ai_settings(ai_max_calls_per_day=1), factory, db_path=tmp_path / "j.db")
    assert len(factory.client.calls) == 1 and run.error.kind == "invalid_response"


def test_disabled_ai_raises_friendly_error(clean, tmp_path):
    with pytest.raises(Exception) as info:
        get_insights(clean, ai_settings(ai_enabled=False), CountingFactory([]), db_path=tmp_path / "x.db")
    assert "turned off" in info.value.user_message


# --- Failures keep everything else working; reports use cached AI -----------------------------

def test_quota_error_logged_and_outputs_still_work(clean, tmp_path):
    db = tmp_path / "q.db"
    factory = CountingFactory([classify_error(quota_error())])
    run = get_insights(clean, ai_settings(), factory, db_path=db)
    assert run.error.kind == "quota" and len(factory.client.calls) == 1
    usage = repo.ai_usage_since(today_start_utc(), db_path=db)
    assert usage["failed"] == 1 and usage["calls"] == 1
    pdf = tmp_path / "r.pdf"
    generate_pdf(clean, pdf, ai=run.insights)
    assert AI_NOT_GENERATED in "".join(p.extract_text() for p in PdfReader(str(pdf)).pages)
    sheets = generate_workbook(clean, tmp_path / "r.xlsx", ai=run.insights)
    assert "AI_Insights" not in sheets


def test_reports_include_cached_ai(clean, tmp_path):
    db = tmp_path / "p.db"
    get_insights(clean, ai_settings(), CountingFactory([answer_for(clean)]), db_path=db)
    cached = get_insights(clean, ai_settings(), None, allow_call=False, db_path=db)
    pdf = tmp_path / "ai.pdf"
    generate_pdf(clean, pdf, ai=cached.insights)
    text = "".join(p.extract_text() for p in PdfReader(str(pdf)).pages)
    assert AI_NOT_GENERATED not in text and "Evidence: E" in text and "HYPOTHESIS" in text
    xlsx = tmp_path / "ai.xlsx"
    assert "AI_Insights" in generate_workbook(clean, xlsx, ai=cached.insights)
    ws = load_workbook(xlsx)["AI_Insights"]
    assert "AI-GENERATED" in ws["A1"].value


# --- App --------------------------------------------------------------------------------------

@pytest.fixture
def app_env(tmp_path, monkeypatch):
    import ai.client
    monkeypatch.setattr(dbc, "settings", Settings(database_path=str(tmp_path / "app.db")))
    monkeypatch.setattr(config_settings, "settings", ai_settings(ai_max_calls_per_day=2))
    holder = {}

    def factory(settings):
        holder.setdefault("client", FakeClient([holder["answer"]] * 5))
        return holder["client"]
    monkeypatch.setattr(ai.client, "client_from_settings", factory)
    return holder


def _app_on_ai_page():
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=240).run()
    at.button(key="load_sample").click().run()
    at.button(key="run_analysis").click().run()
    at.sidebar.radio(key="page").set_value("AI Insights").run()
    return at


def test_app_cache_regenerate_confirmation_and_usage(app_env, clean):
    app_env["answer"] = answer_for(clean)
    at = _app_on_ai_page()
    at.button(key="generate_ai").click().run()
    assert not at.exception and len(app_env["client"].calls) == 1

    # A new session on the same database: saved insights load with no call.
    at2 = _app_on_ai_page()
    assert any("Loaded saved insights" in s.value for s in at2.success)
    assert len(app_env["client"].calls) == 1

    # Regenerate needs a confirmation click.
    at2.button(key="regenerate_ai").click().run()
    assert len(app_env["client"].calls) == 1
    at2.button(key="confirm_regenerate").click().run()
    assert not at2.exception and len(app_env["client"].calls) == 2

    # Daily limit (2) now reached: the regenerate confirmation is disabled.
    at2.button(key="regenerate_ai").click().run()
    assert at2.button(key="confirm_regenerate").disabled
    assert any("daily AI limit" in c.value for c in at2.caption)

    at2.sidebar.radio(key="page").set_value("Data Quality").run()
    metrics = {m.label: m.value for m in at2.metric}
    assert metrics["AI calls today"] == "2 of 2" and metrics["Calls left now"] == "0"
    assert metrics["Saved results reused today"] == "0"     # page-open loads are free and not logged

    at2.sidebar.radio(key="page").set_value("Reports").run()
    at2.button(key="generate_pdf").click().run()
    assert not at2.exception and any("Report ready" in s.value for s in at2.success)


def test_app_ai_off_message(tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    monkeypatch.setattr(dbc, "settings", Settings(database_path=str(tmp_path / "app.db")))
    monkeypatch.setattr(config_settings, "settings", Settings(ai_enabled=False))
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=240).run()
    at.button(key="load_sample").click().run()
    at.button(key="run_analysis").click().run()
    at.sidebar.radio(key="page").set_value("AI Insights").run()
    assert any("AI is turned off" in i.value for i in at.info)
    at.sidebar.radio(key="page").set_value("Data Quality").run()
    assert not at.exception
