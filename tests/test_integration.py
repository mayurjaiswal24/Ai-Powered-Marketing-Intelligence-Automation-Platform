"""Phase 13: end-to-end flow, cross-output consistency and reopening saved runs."""

import json
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

import database.connection as dbc
import reports.excel_report as excel_report
import reports.pdf_report as pdf_report
from ai.cache import get_insights
from analytics.channel import channel_table
from analytics.kpis import compute_kpis
from config.settings import Settings
from dashboard.pipeline import SAMPLE_DIR, load_raw, prepare, reopen_run, run_pipeline
from dashboard.tables import format_table
from database import repository as repo
from database.connection import connect
from reports.excel_report import export_workbook
from reports.pdf_report import AI_NOT_GENERATED, export_pdf
from tests.test_ai import FakeClient, good_answer
from utils.formatting import title_case

ROOT = Path(__file__).resolve().parent.parent
CHANNEL = "Paid Search"


def pdf_text(path) -> str:
    return " ".join("".join(p.extract_text() for p in PdfReader(str(path)).pages).split())


@pytest.fixture(scope="module")
def flow(tmp_path_factory):
    """The full workflow once, AI off: sample -> profile -> mapping -> clean -> store -> analyse
    -> PDF -> Excel."""
    folder = tmp_path_factory.mktemp("e2e")
    db = folder / "app.db"
    raw, report = load_raw(SAMPLE_DIR / "marketing_clean.csv")
    prep = prepare(raw, report)
    assert prep.validation.can_analyse and not prep.mapping.needs_confirmation
    out = run_pipeline(prep, db_path=db)
    pdf = export_pdf(out.analysis, exports_dir=folder / "exports", db_path=db)
    xlsx = export_workbook(out.analysis, out.clean.clean_df, out.clean.quality_log_df,
                           exports_dir=folder / "exports", db_path=db)
    return {"db": db, "out": out, "pdf": pdf, "xlsx": xlsx, "folder": folder}


# --- End to end ---------------------------------------------------------------------------------

def test_end_to_end_without_ai(flow):
    out, db = flow["out"], flow["db"]
    assert out.save_error is None and out.run_id is not None
    run = repo.get_run(out.run_id, db_path=db)
    assert run["status"] == "complete" and run["file_name"] == "marketing_clean.csv"
    assert len(repo.load_clean_records(out.run_id, db_path=db)) == 8471
    assert not repo.load_quality_log(out.run_id, db_path=db).empty or out.clean.quality_log_df.empty
    assert len(repo.load_analysis_table(out.run_id, "anomalies", db_path=db)) == len(out.analysis.anomalies)
    reports = repo.list_reports(out.run_id, db_path=db)
    assert sorted(reports["report_type"]) == ["excel", "pdf"]
    for path in (flow["pdf"], flow["xlsx"]):
        assert path.exists() and path.parent.name == str(out.run_id)
        assert re.search(rf"_run{out.run_id}_\d{{4}}-\d{{2}}-\d{{2}}_\d{{6}}\.", path.name)
    assert AI_NOT_GENERATED in pdf_text(flow["pdf"])
    assert "AI_Insights" not in load_workbook(flow["xlsx"]).sheetnames


def test_end_to_end_with_fake_ai(flow):
    out, db = flow["out"], flow["db"]
    settings = Settings(ai_enabled=True, gemini_api_key="k", gemini_model="fake-model")
    from ai.context_builder import build_evidence
    client = FakeClient([json.dumps(good_answer(build_evidence(out.analysis, model="fake-model")))])
    run = get_insights(out.analysis, settings, lambda: client, run_id=out.run_id, db_path=db)
    assert run.ok and len(client.calls) == 1
    with connect(db) as conn:
        ai_row = conn.execute("SELECT run_id, model, success, cache_hit FROM ai_runs").fetchone()
        stored = conn.execute("SELECT run_id, created_at FROM ai_insights").fetchone()
    assert ai_row == (out.run_id, "fake-model", 1, 0) and stored[0] == out.run_id and stored[1]
    pdf = export_pdf(out.analysis, exports_dir=flow["folder"] / "ai", db_path=db, ai=run.insights)
    xlsx = export_workbook(out.analysis, out.clean.clean_df, out.clean.quality_log_df,
                           exports_dir=flow["folder"] / "ai", db_path=db, ai=run.insights)
    assert AI_NOT_GENERATED not in pdf_text(pdf) and "Evidence: E" in pdf_text(pdf)
    assert "AI_Insights" in load_workbook(xlsx).sheetnames


# --- One source of truth ------------------------------------------------------------------------

def test_cross_output_consistency(flow):
    """Overall ROAS, CPL, CAC, revenue and one channel's CPL are identical everywhere."""
    out = flow["out"]
    result = out.analysis
    df = out.clean.clean_df
    k = result.kpis

    # 1. Dashboard helpers (the functions the dashboard calls for an unfiltered view).
    dash = compute_kpis(df, result.capabilities)
    for key in ("roas", "cpl", "cac", "revenue"):
        assert dash[key].value == k[key].value
        assert dash[key].formatted_compact == k[key].formatted_compact     # KPI card text
    dash_channels = channel_table(df).set_index("channel")
    ch_cpl = result.channels.set_index("channel").loc[CHANNEL, "cpl"]
    assert dash_channels.loc[CHANNEL, "cpl"] == ch_cpl
    shown = format_table(dash_channels.reset_index(), ["channel", "cpl"]).set_index("Channel")
    cpl_text = shown.loc[CHANNEL, "CPL"]

    # 2. PDF text.
    text = pdf_text(flow["pdf"])
    for key in ("roas", "cpl", "cac", "revenue"):
        assert k[key].formatted in text, key           # KPI overview (full format)
    assert f"{CHANNEL} " in text and cpl_text in text

    # 3. Excel cells (real numbers).
    wb = load_workbook(flow["xlsx"])
    ws = wb["Executive_KPIs"]
    values = {r[0].value: r[1].value for r in ws.iter_rows(min_row=12, max_col=2)}
    for key in ("roas", "cpl", "cac", "revenue"):
        assert values[title_case(k[key].label)] == pytest.approx(k[key].value, rel=1e-12)
    cws = wb["Channel_Analysis"]
    headers = [c.value for c in cws[3]]
    row = next(r for r in cws.iter_rows(min_row=4) if r[0].value == CHANNEL)
    assert row[headers.index("CPL")].value == pytest.approx(ch_cpl, rel=1e-12)

    # 4. Database.
    stored = repo.load_analysis_table(out.run_id, "kpi_results", db_path=flow["db"]).set_index("kpi")
    for key in ("roas", "cpl", "cac", "revenue"):
        assert stored.loc[key, "value"] == pytest.approx(k[key].value, rel=1e-12)


# --- Recent analyses ----------------------------------------------------------------------------

def test_reopen_run_without_recomputing(flow):
    out, db = flow["out"], flow["db"]
    recent = repo.list_recent_runs(db_path=db)
    assert recent.iloc[0]["run_id"] == out.run_id and recent.iloc[0]["reports"] >= 2
    again = reopen_run(out.run_id, db_path=db)
    assert again.analysis.kpis["roas"].value == out.analysis.kpis["roas"].value
    assert len(again.analysis.incidents) == len(out.analysis.incidents)
    assert again.clean.clean_df.equals(out.clean.clean_df)


def test_reopen_unknown_or_broken_snapshot(flow):
    with pytest.raises(dbc.DatabaseError) as info:
        reopen_run(99999, db_path=flow["db"])
    assert "upload the file again" in info.value.user_message
    repo.save_run_snapshot(flow["out"].run_id, b"not a pickle", db_path=flow["db"])
    with pytest.raises(dbc.DatabaseError) as info:
        reopen_run(flow["out"].run_id, db_path=flow["db"])
    assert "older version" in info.value.user_message


def test_app_recent_analyses_reopen(tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    monkeypatch.setattr(dbc, "settings", Settings(database_path=str(tmp_path / "app.db")))
    monkeypatch.setattr(pdf_report, "EXPORTS_DIR", tmp_path / "exports")
    monkeypatch.setattr(excel_report, "EXPORTS_DIR", tmp_path / "exports")
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=240).run()
    at.button(key="load_sample").click().run()
    at.button(key="run_analysis").click().run()
    # A new browser session: reopen the saved run from "Recent analyses".
    at2 = AppTest.from_file(str(ROOT / "app.py"), default_timeout=240).run()
    at2.button(key="open_run_1").click().run()
    assert not at2.exception
    assert any("Opened Analysis #1 " in s.value for s in at2.success)
    for page in ("Executive Overview", "Channels", "Anomalies", "Data Quality", "Reports"):
        at2.sidebar.radio(key="page").set_value(page).run()
        assert not at2.exception and not at2.error, page
    at2.button(key="generate_excel").click().run()
    assert any("workbook is ready" in s.value for s in at2.success)
