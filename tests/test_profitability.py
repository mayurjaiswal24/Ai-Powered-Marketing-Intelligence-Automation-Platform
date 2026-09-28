"""U4 Profitability: hand-calculated fixture, status boundaries, assumed margin, no revenue,
owned channels, consistency with ROI, exports and the app page."""

import io
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

import database.connection as dbc
from analytics.engine import run_analysis
from analytics.kpis import (LOSS, NEAR, NOT_AVAILABLE, PROFITABLE, compute_kpis, profit_status)
from analytics.profitability import (REASON_NO_MARGIN, REASON_NO_REVENUE, assumed_margin_label,
                                     profitability_analysis, summary_lines)
from config.settings import PROFIT_NEAR_BAND, Settings
from dashboard.pipeline import load_raw, prepare, run_pipeline
from reports.excel_report import generate_workbook
from reports.pdf_report import generate_pdf

ROOT = Path(__file__).resolve().parent.parent
SAMPLE_DIR = ROOT / "data" / "sample"

# Each campaign is split over two days, so every figure must come from summed totals.
#   A: spend 100, revenue 400, GP 200 -> margin 50%, break-even 2.00x, ROAS 4.00x, contribution 100
#   B: spend 100, revenue 220, GP 110 -> ROAS 2.20x = exactly 1.10 x break-even -> Profitable
#   C: spend 100, revenue 180, GP  90 -> ROAS 1.80x = exactly 0.90 x break-even -> Near break-even
#   D: spend 100, revenue 150, GP  75 -> ROAS 1.50x = 0.75 x break-even -> Loss-making (-25)
#   E: Email (owned): spend 10, revenue 100, GP 50 -> reported separately, never rated
ROWS = [  # campaign, channel, spend, revenue, gross profit (day 1 and day 2)
    ("A", "Paid Search", (40, 60), (150, 250), (100, 100)),
    ("B", "Paid Social", (50, 50), (100, 120), (50, 60)),
    ("C", "Paid Social", (30, 70), (80, 100), (40, 50)),
    ("D", "Video", (60, 40), (90, 60), (45, 30)),
    ("E", "Email", (5, 5), (40, 60), (20, 30)),
]


def fixture_df() -> pd.DataFrame:
    records = []
    for name, channel, spend, revenue, gp in ROWS:
        for day in range(2):
            date = pd.Timestamp("2026-01-30") + pd.Timedelta(days=day * 3)     # January and February
            records.append({"date": date, "month": date.strftime("%Y-%m"), "campaign_name": name,
                            "channel": channel, "spend": float(spend[day]), "revenue": float(revenue[day]),
                            "gross_profit": float(gp[day])})
    return pd.DataFrame(records)


def test_hand_calculated_fixture():
    prof = profitability_analysis(fixture_df())
    assert prof.available and not prof.assumed
    camps = prof.campaigns.set_index("campaign")
    a = camps.loc["A"]
    assert a["gross_margin_pct"] == pytest.approx(50.0)
    assert a["break_even_roas"] == pytest.approx(2.0)
    assert a["roas"] == pytest.approx(4.0)
    assert a["contribution"] == pytest.approx(100.0)
    assert a["profit_per_rupee"] == pytest.approx(1.0)
    assert a["roas_headroom"] == pytest.approx(100.0)
    assert camps.loc["D", "contribution"] == pytest.approx(-25.0)
    assert camps.loc["D", "roas_headroom"] == pytest.approx(-25.0)
    assert list(camps["status"].reindex(list("ABCD"))) == [PROFITABLE, PROFITABLE, NEAR, LOSS]
    assert prof.status_counts == {PROFITABLE: 2, NEAR: 1, LOSS: 1, NOT_AVAILABLE: 0}
    assert prof.loss_spend == pytest.approx(100.0)
    assert prof.loss_spend_share_pct == pytest.approx(25.0)
    assert list(prof.loss_makers()["campaign"]) == ["D"]
    # Channels from their own totals: Paid Social = B + C (spend 200, revenue 400, GP 200).
    social = prof.channels.set_index("channel").loc["Paid Social"]
    assert social["roas"] == pytest.approx(2.0) and social["contribution"] == pytest.approx(0.0)
    assert social["status"] == NEAR
    # Months: 30 Jan (Jan) and 2 Feb (Feb).
    assert list(prof.monthly["period"]) == ["2026-01", "2026-02"]
    assert prof.monthly["contribution"].sum() == pytest.approx(prof.totals["contribution"])


def test_owned_channels_are_excluded_from_ratings():
    prof = profitability_analysis(fixture_df())
    assert "E" not in set(prof.campaigns["campaign"])
    assert list(prof.owned_campaigns["campaign"]) == ["E"]
    assert "Email" not in set(prof.channels["channel"])
    assert list(prof.owned_channels["channel"]) == ["Email"]
    assert sum(prof.status_counts.values()) == 4                 # paid campaigns only
    # ... but the overall totals include them (the same rows as ROI).
    assert prof.totals["spend"] == pytest.approx(410.0)


def test_overall_contribution_is_total_gp_minus_total_spend_consistent_with_roi():
    df = fixture_df()
    prof = profitability_analysis(df)
    assert prof.totals["contribution"] == pytest.approx(525.0 - 410.0)
    kpis = compute_kpis(df)
    assert prof.totals["contribution"] == pytest.approx(kpis["roi"].numerator)
    assert prof.totals["profit_per_rupee"] * 100 == pytest.approx(kpis["roi"].value)
    assert prof.totals["roas"] == pytest.approx(kpis["roas"].value)


@pytest.mark.parametrize("roas,expected", [
    (2.2, PROFITABLE),            # exactly 1.10 x break-even
    (2.1998, NEAR),               # just under 1.10
    (2.0, NEAR),
    (1.8, NEAR),                  # exactly 0.90 x break-even
    (1.7998, LOSS),               # just under 0.90
    (0.0, LOSS),
])
def test_status_boundaries(roas, expected):
    status = profit_status(pd.Series([roas]), pd.Series([2.0]), band=0.10)
    assert status.iloc[0] == expected


def test_status_when_break_even_is_missing():
    status = profit_status(pd.Series([np.nan, 1.0, np.nan]), pd.Series([np.nan, np.nan, np.nan]),
                           pd.Series([-50.0, -10.0, 20.0]))
    assert list(status) == [LOSS, LOSS, NOT_AVAILABLE]
    assert PROFIT_NEAR_BAND == 0.10


def test_assumed_margin_path():
    df = fixture_df().drop(columns="gross_profit")
    blank = profitability_analysis(df)
    assert not blank.available and blank.needs_margin and blank.reason == REASON_NO_MARGIN
    prof = profitability_analysis(df, 50)
    assert prof.available and prof.assumed and prof.assumed_margin_pct == 50
    assert prof.basis_note.startswith("Based on your assumed margin of 50%")
    assert assumed_margin_label(37.5) == "Based on your assumed margin of 37.5%"
    # GP = revenue x 50%: A has revenue 400 -> GP 200, the same as the fixture's real figure.
    a = prof.campaigns.set_index("campaign").loc["A"]
    assert a["gross_profit"] == pytest.approx(200.0) and a["break_even_roas"] == pytest.approx(2.0)
    # A margin in the data always wins over an assumed one.
    real = profitability_analysis(fixture_df(), 10)
    assert not real.assumed and real.basis_note == "Based on the gross profit column in the data."


def test_margin_column_is_used_before_any_assumption():
    df = fixture_df().drop(columns="gross_profit").assign(margin=50.0)
    prof = profitability_analysis(df, 10)
    assert not prof.assumed and "margin column" in prof.basis_note
    assert prof.totals["gross_margin_pct"] == pytest.approx(50.0)


def test_no_revenue_explains_instead_of_failing():
    df = fixture_df().drop(columns=["revenue", "gross_profit"])
    prof = profitability_analysis(df, 40)
    assert not prof.available and not prof.needs_margin and prof.reason == REASON_NO_REVENUE
    assert summary_lines(prof) == [REASON_NO_REVENUE]


def test_engine_adds_profitability_without_changing_tables():
    df = fixture_df()
    df["week_start"] = df["date"] - pd.to_timedelta(df["date"].dt.weekday, unit="D")
    result = run_analysis(df)
    assert result.profitability.available
    assert "profitability" not in result.tables()


@pytest.fixture(scope="module")
def kalpa(tmp_path_factory):
    raw, report = load_raw(SAMPLE_DIR / "marketing_clean.csv")
    return run_pipeline(prepare(raw, report), db_path=tmp_path_factory.mktemp("db") / "k.db")


def test_kalpa_profitability_matches_kpis(kalpa):
    a = kalpa.analysis
    prof = a.profitability
    assert prof.available and not prof.assumed
    assert prof.totals["contribution"] == pytest.approx(a.kpis["roi"].numerator)
    assert prof.totals["roas"] == pytest.approx(a.kpis["roas"].value)
    campaign_roas = a.campaigns.set_index("campaign")["roas"]
    for row in prof.campaigns.itertuples():
        assert row.roas == pytest.approx(campaign_roas[row.campaign])
    assert "Email" in set(prof.owned_channels["channel"])


def test_exports_label_the_assumed_margin(tmp_path):
    raw, report = load_raw(SAMPLE_DIR / "marketing_no_margin.csv")
    out = run_pipeline(prepare(raw, report), db_path=tmp_path / "m.db")
    assert not out.analysis.profitability.available
    prof = profitability_analysis(out.clean.clean_df, 35)
    path = tmp_path / "m.xlsx"
    sheets = generate_workbook(out.analysis, path, profitability=prof)
    assert "Profitability" in sheets
    ws = load_workbook(path)["Profitability"]
    assert ws["A1"].value.startswith("Based on your assumed margin of 35%")
    buffer = io.BytesIO()
    generate_pdf(out.analysis, buffer, profitability=prof)
    text = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(buffer.getvalue())).pages)
    assert "Based on your assumed margin of 35%" in text
    # Without an assumed margin there is no profitability sheet.
    assert "Profitability" not in generate_workbook(out.analysis, tmp_path / "n.xlsx")


def test_profitability_page_with_assumed_margin(kalpa, tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    monkeypatch.setattr(dbc, "settings", Settings(database_path=str(tmp_path / "app.db")))
    raw, report = load_raw(SAMPLE_DIR / "marketing_no_margin.csv")
    no_margin = run_pipeline(prepare(raw, report), db_path=tmp_path / "m.db")
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=240).run()
    at.session_state["output"] = no_margin
    at.sidebar.radio(key="page").set_value("Profitability").run()
    assert not at.exception
    assert at.number_input(key="assumed_margin_pct").value is None       # blank by default
    at.number_input(key="assumed_margin_pct").set_value(40.0).run()
    assert not at.exception
    assert any("Based on your assumed margin of 40%" in w.value for w in at.warning)
