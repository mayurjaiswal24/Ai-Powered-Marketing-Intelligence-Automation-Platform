"""Phase 14 changes: random cross-check sampling, paid-media ROAS in the evidence, and
"₹ at stake" for recommendations (calculated in code, used for ranking). FAKE AI only."""

import numpy as np
import pandas as pd
import pytest

from ai.context_builder import build_evidence
from ai.evaluator import evaluate
from ai.prompts import PROMPT_VERSION, SYSTEM_INSTRUCTIONS
from ai.schemas import AIInsights
from analytics.channel import paid_media_kpis
from analytics.kpis import compute_kpis
from dashboard.pipeline import SAMPLE_DIR, load_raw, prepare, run_pipeline
from ingestion.crosscheck import cross_check, sample_rows
from ingestion.mapper import map_fields
from tests.test_ai import good_answer


@pytest.fixture(scope="module")
def analysis(tmp_path_factory):
    raw, report = load_raw(SAMPLE_DIR / "marketing_clean.csv")
    return run_pipeline(prepare(raw, report), db_path=tmp_path_factory.mktemp("p14") / "a.db").analysis


@pytest.fixture(scope="module")
def pack(analysis):
    return build_evidence(analysis, model="m")


# --- 2. Cross-check sampling ---------------------------------------------------------------------

def test_sample_spreads_over_the_whole_file_and_is_repeatable():
    df = pd.DataFrame({"i": range(5000)})
    first, second = sample_rows(df), sample_rows(df)
    assert len(first) == 500 and first.index.equals(second.index)          # fixed seed
    assert first["i"].min() < 500 and first["i"].max() > 4500               # not just the top
    assert first.index.is_monotonic_increasing                              # order kept
    small = pd.DataFrame({"i": range(40)})
    assert sample_rows(small).equals(small)                                 # small file: all rows


def test_cross_check_uses_rows_beyond_the_top():
    rng = np.random.default_rng(3)
    n = 2000
    spend = rng.uniform(1000, 5000, n).round(2)
    clicks = rng.integers(100, 900, n)
    cpc = [str(v) for v in (spend / clicks).round(2)]
    cpc[:600] = [""] * 600                              # the file's own CPC starts at row 601
    df = pd.DataFrame({"Spend": spend.astype(str), "Clicks": clicks.astype(str), "CPC": cpc})
    checks = {c.column: c.status for c in cross_check(df, map_fields(df)).checks}
    assert checks["CPC"] == "verified"                  # the first 500 rows alone could not tell


# --- 3. Paid-media ROAS in the evidence ----------------------------------------------------------

def test_paid_media_roas_is_its_own_figure(analysis, pack):
    kpi = next(i for i in pack.items if i.type == "kpi")
    assert "ROAS" not in kpi.facts
    assert kpi.facts["ROAS overall (all channels)"] == analysis.kpis["roas"].formatted_compact
    paid = analysis.paid_kpis["roas"]
    assert kpi.facts["ROAS paid media only (owned channels excluded)"] == paid.formatted_compact
    assert paid.value != analysis.kpis["roas"].value                      # email is excluded


def test_paid_media_kpis_come_from_paid_rows_only():
    df = pd.DataFrame({"channel": ["Paid Search", "Email", "Paid Search"],
                       "spend": [100.0, 10.0, 100.0], "revenue": [500.0, 400.0, 300.0]})
    assert paid_media_kpis(df)["roas"].value == pytest.approx(800 / 200)
    assert compute_kpis(df)["roas"].value == pytest.approx(1200 / 210)
    assert paid_media_kpis(df.drop(columns="channel")) == {}


# --- 4. Rupees at stake ----------------------------------------------------------------------------

def _items(pack, kind):
    return [i for i in pack.items if i.type == kind]


def _rec(text, ids, **extra):
    return {"text": text, "evidence_ids": ids, "priority": "medium", "metric_to_watch": "ROAS", **extra}


def test_at_stake_is_test_share_times_source_spend_and_ranks_everything(pack):
    campaigns = [i for i in _items(pack, "campaign") if "spend" in i.values]
    channels = [i for i in _items(pack, "channel") if "spend" in i.values]
    incident = _items(pack, "incident")[0]
    source, receiver = channels[0], channels[-1]
    small = min(campaigns, key=lambda c: c.values["spend"])
    impact = incident.facts["estimated impact"].split(" (")[0].rsplit(" ", 1)[0]
    answer = good_answer(pack)
    answer["recommendations"] = [
        _rec("No money involved here.", [pack.items[0].id]),
        _rec("Small test on a small campaign.", [small.id], test_shift_pct=10, source_evidence_id=small.id),
        _rec("Move a test share from the big channel.", [source.id, receiver.id], test_shift_pct=20,
             source_evidence_id=source.id),
        _rec(f"Prevent a repeat: {impact}.", [incident.id]),
    ]
    checked, _ = evaluate(AIInsights.model_validate(answer).model_dump(), pack)
    recs = checked["recommendations"]
    by_text = {r["text"]: r for r in recs}
    big = by_text["Move a test share from the big channel."]
    assert big["rupees_at_stake"] == pytest.approx(0.20 * source.values["spend"])
    assert big["at_stake_text"].startswith("₹ at stake: ₹") and "not a projected gain" in big["at_stake_text"]
    assert by_text["Small test on a small campaign."]["rupees_at_stake"] == pytest.approx(
        0.10 * small.values["spend"])
    none = by_text["No money involved here."]
    assert none["rupees_at_stake"] == 0 and not none["at_stake_text"]
    assert by_text[f"Prevent a repeat: {impact}."]["rupees_at_stake"] > 0     # incident impact
    amounts = [r["rupees_at_stake"] for r in recs]
    assert amounts == sorted(amounts, reverse=True)                      # ranked by ₹ at stake
    assert recs[-1]["text"] == "No money involved here."


def test_source_is_inferred_for_older_answers(pack):
    channels = sorted((i for i in _items(pack, "channel") if "roas" in i.values),
                      key=lambda c: c.values["roas"])
    weakest, strongest = channels[0], channels[-1]
    answer = good_answer(pack)
    answer["recommendations"][0] = _rec("Test a move.", [strongest.id, weakest.id], test_shift_pct=15)
    checked, _ = evaluate(AIInsights.model_validate(answer).model_dump(), pack)
    rec = next(r for r in checked["recommendations"] if r["text"] == "Test a move.")
    assert rec["rupees_at_stake"] == pytest.approx(0.15 * weakest.values["spend"])   # from the weaker


def test_exact_values_are_not_sent_to_gemini(pack):
    assert '"values"' not in pack.json_text


def test_prompt_v5_rules():
    assert PROMPT_VERSION >= "v5"             # the v5 rules stay in later versions
    text = " ".join(SYSTEM_INSTRUCTIONS.lower().split())
    for rule in ("source_evidence_id", "never state a projected gain", "roas paid media only"):
        assert rule in text, rule


def test_excel_ai_sheet_has_rupees_at_stake(tmp_path, analysis, pack):
    from openpyxl import load_workbook
    from reports.excel_report import export_workbook
    source = next(i for i in _items(pack, "channel") if "spend" in i.values)
    answer = good_answer(pack)
    answer["recommendations"][0] = _rec("Move a test share.", [source.id], test_shift_pct=10,
                                        source_evidence_id=source.id)
    checked, _ = evaluate(AIInsights.model_validate(answer).model_dump(), pack)
    path = export_workbook(analysis, exports_dir=tmp_path, db_path=tmp_path / "x.db", ai=checked)
    rows = list(load_workbook(path)["AI_Insights"].iter_rows(values_only=True))
    head = next(i for i, r in enumerate(rows) if any("at stake" in str(v).lower() for v in r if v))
    col = next(j for j, v in enumerate(rows[head]) if v and "at stake" in str(v).lower())
    values = [r[col] for r in rows[head + 1:] if isinstance(r[col], (int, float))]
    assert values and values[0] == pytest.approx(0.10 * source.values["spend"])   # a real number
