"""Builds the small, deliberately awkward files used by tests/test_loader.py.

Run from the project root:  python tests/fixtures/build_fixtures.py
The generated files are committed, so this only needs re-running if a fixture changes.
"""

import zipfile
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).resolve().parent

HEADER = "date,campaign_name,channel,spend,clicks,leads\n"
ROWS = [
    "2026-01-05,Google Search - Data,Paid Search,1200.50,40,4\n",
    "2026-01-06,Google Search - Data,Paid Search,1350.00,45,5\n",
    "2026-01-07,Meta - Prospecting,Paid Social,800.25,90,3\n",
]


def write_bytes(name: str, data: bytes) -> None:
    (HERE / name).write_bytes(data)


def build() -> None:
    write_bytes("empty.csv", b"")
    write_bytes("header_only.csv", HEADER.encode())
    write_bytes("fake.xlsx", b"This is just a text file that someone renamed to .xlsx\n")
    write_bytes("report.pdf", b"%PDF-1.4 not really a pdf\n")
    write_bytes("legacy.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504)

    # Semicolon-separated with European decimal commas (common from European Excel exports).
    write_bytes("semicolon.csv", (
        "date;campaign_name;channel;spend;clicks;leads\n"
        "2026-01-05;Google Search - Data;Paid Search;1200,50;40;4\n"
        "2026-01-06;Google Search - Data;Paid Search;1350,00;45;5\n"
        "2026-01-07;Meta - Prospecting;Paid Social;800,25;90;3\n").encode())

    write_bytes("tab_delimited.csv", (HEADER + "".join(ROWS)).replace(",", "\t").encode())

    # Latin-1: "é" is a single byte 0xE9, which is invalid UTF-8.
    write_bytes("latin1.csv", (
        "date,campaign_name,channel,spend\n"
        "2026-01-05,Café Learners Campaign,Paid Social,500\n"
        "2026-01-06,Résumé Builders,Paid Search,700\n").encode("latin-1"))

    # UTF-8 with a byte-order mark (Excel's "CSV UTF-8" option) and a rupee sign.
    write_bytes("utf8_bom.csv", ("date,campaign_name,spend\n"
                                 '2026-01-05,Brand,"₹1,23,456"\n').encode("utf-8-sig"))

    # Structural oddities: duplicate and blank headings, an empty row in the middle.
    write_bytes("odd_structure.csv", (
        "date,spend,spend,,leads\n"
        "2026-01-05,100,110,x,1\n"
        ",,,,\n"
        "2026-01-06,200,210,y,2\n").encode())

    # Two sheets: a small notes sheet first, the real data second.
    wb = Workbook()
    notes = wb.active
    notes.title = "Notes"
    notes["A1"] = "Exported from the ad platform"
    notes["A2"] = "Owner: marketing team"
    data = wb.create_sheet("Campaign Data")
    data.append(["date", "campaign_name", "channel", "spend", "clicks", "leads"])
    for row in ROWS:
        d, name, channel, spend, clicks, leads = row.strip().split(",")
        data.append([d, name, channel, float(spend), int(clicks), int(leads)])
    wb.save(HERE / "two_sheets.xlsx")

    # A report title and a blank row above the real header.
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "Kalpa Learning - Campaign Performance Report"
    ws.append([])
    ws.append(["date", "campaign_name", "channel", "spend", "clicks", "leads"])
    for row in ROWS:
        d, name, channel, spend, clicks, leads = row.strip().split(",")
        ws.append([d, name, channel, float(spend), int(clicks), int(leads)])
    wb.save(HERE / "title_row.xlsx")

    # Stand-in for a password-protected workbook: Excel saves those as an OLE container with an
    # "EncryptedPackage" stream. Creating a real one needs extra tools, so this file reproduces
    # the two markers the loader checks for.
    write_bytes("password_protected.xlsx",
                b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504
                + "EncryptedPackage".encode("utf-16-le") + b"\x00" * 64)

    # A valid ZIP file that is not a workbook.
    with zipfile.ZipFile(HERE / "not_a_workbook.xlsx", "w") as zf:
        zf.writestr("readme.txt", "hello")

    # A workbook whose only sheet is empty.
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.save(HERE / "empty_workbook.xlsx")


if __name__ == "__main__":
    build()
    print(f"Fixtures written to {HERE}")
