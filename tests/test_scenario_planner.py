"""U7 Scenario planner: recovering a known curve, clamping, the data check (insufficient / low
variation), incident weeks left out, anchoring, performance-priced channels never fitted (with
their cap), total spend conserved, guardrails, marginal cost vs a numeric derivative, a
deterministic bootstrap, Suggest Allocation limits, exports, the page and the AI pre-fill. The AI
evidence pack is unchanged (tests/test_demo.py pins its fingerprint)."""

import io
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

import ai.demo_seed
import dashboard.demo
import database.connection as dbc
from analytics import response_curves as rc
from analytics.kpis import marginal_cost_per_conversion, response_conversions, response_scale
from config import settings as cfg
from config.settings import Settings
from dashboard.pipeline import load_raw, prepare, run_pipeline
from reports.excel_report import generate_workbook, scenario_workbook
from reports.pdf_report import generate_pdf

ROOT = Path(__file__).resolve().parent.parent
SAMPLE_DIR = ROOT / "data" / "sample"


def weekly(spend, conversions, start="2025-01-06") -> pd.DataFrame:
    periods = pd.date_range(start, periods=len(spend), freq="7D")
    return pd.DataFrame({"period": periods, "spend": np.asarray(spend, float),
                         "conversions": np.asarray(conversions, float),
                         "revenue": np.asarray(conversions, float) * 1000.0})


def power_data(b=0.6, a=2.0, weeks=40, noise=0.05, seed=1):
    rng = np.random.default_rng(seed)
    spend = np.linspace(50_000, 200_000, weeks)[rng.permutation(weeks)]
    conv = a * spend ** b * np.exp(rng.normal(0, noise, weeks))
    return spend, conv


def daily_rows(channel_weeks: dict, start="2025-01-06") -> pd.DataFrame:
    """Daily rows (7 per week) per channel from weekly spend / conversions lists."""
    rows = []
    for channel, (spend, conv) in channel_weeks.items():
        for w, (s, c) in enumerate(zip(spend, conv)):
            for d in range(7):
                rows.append({"date": pd.Timestamp(start) + pd.Timedelta(days=7 * w + d), "channel": channel,
                             "campaign_name": f"{channel} A", "spend": s / 7, "conversions": c / 7,
                             "revenue": c / 7 * 1000.0})
    return pd.DataFrame(rows)


# --- Formulas ---------------------------------------------------------------------------------------

def test_formulas_and_marginal_cost_vs_numeric_derivative():
    a, b, s = 2.0, 0.6, 120_000.0
    assert response_conversions(s, a, b) == pytest.approx(a * s ** b)
    assert response_conversions(0, a, b) == 0.0
    h = 1.0
    slope = (response_conversions(s + h, a, b) - response_conversions(s - h, a, b)) / (2 * h)
    assert marginal_cost_per_conversion(s, a, b) == pytest.approx(1 / slope, rel=1e-6)
    assert marginal_cost_per_conversion(0, a, b) is None
    assert response_scale(100, 10_000, 0.5) == pytest.approx(1.0)


def test_recovers_b_on_synthetic_data():
    spend, conv = power_data(b=0.6)
    b, a, r2 = rc.fit_power(spend, conv)
    assert b == pytest.approx(0.6, abs=0.05) and r2 > 0.9
    curve = rc.fit_channel(weekly(spend, conv), "Paid Search")
    assert curve.status == rc.FITTED and curve.b == pytest.approx(0.6, abs=0.05)
    assert curve.confidence == "High" and not curve.clamped


def test_clamping_and_no_positive_response():
    spend = np.linspace(50_000, 200_000, 30)
    steep = rc.fit_channel(weekly(spend, 1e-6 * spend ** 1.5), "Paid Search")
    assert steep.b_raw > 1 and steep.b == 1.0 and steep.clamped and steep.confidence == "Low"
    assert "limited" in rc.clamp_note(steep)
    falling = rc.fit_channel(weekly(spend, 1e7 / spend), "Paid Search")
    assert falling.status == rc.INSUFFICIENT and "Insufficient evidence" in falling.reason and not falling.usable


def test_anchor_passes_through_recent_average():
    spend, conv = power_data(b=0.6, weeks=40)
    curve = rc.fit_channel(weekly(spend, conv), "Video")
    recent_s, recent_c = spend[-8:].mean(), conv[-8:].mean()
    assert curve.base_spend == pytest.approx(recent_s)
    assert curve.conversions(recent_s) == pytest.approx(recent_c)
    low, high = curve.conversions_range(recent_s)
    assert low == pytest.approx(recent_c) and high == pytest.approx(recent_c)   # every refit passes the anchor


def test_bootstrap_is_deterministic_and_gives_a_range():
    spend, conv = power_data(b=0.6, noise=0.2)
    one, two = rc.bootstrap_b(spend, conv), rc.bootstrap_b(spend, conv)
    assert len(one) == cfg.RESPONSE_BOOTSTRAP_SAMPLES and np.array_equal(one, two)
    assert ((one >= 0) & (one <= 1)).all()
    curve = rc.fit_channel(weekly(spend, conv), "Video")
    low, high = curve.conversions_range(curve.base_spend * 1.3)
    assert low < curve.conversions(curve.base_spend * 1.3) < high


# --- Data check -----------------------------------------------------------------------------------

def test_data_check_insufficient_and_low_variation():
    flat = np.full(30, 100_000.0) * (1 + 0.02 * np.sin(np.arange(30)))
    v = rc.spend_variation(flat)
    assert v["weeks_with_spend"] == 30 and v["cv"] < cfg.RESPONSE_MIN_CV and not rc.passes_check(v)
    short = rc.spend_variation(np.linspace(50_000, 200_000, 15))
    assert short["weeks_with_spend"] == 15 and not rc.passes_check(short)
    spend, conv = power_data(weeks=30)
    df = daily_rows({"Paid Search": (spend, conv), "Video": (flat, flat / 1000), "Email": (spend, conv)})
    result = rc.response_analysis(df)
    checks = result.checks.set_index("channel")
    assert "Email" not in checks.index                                     # owned channels left out
    assert checks.loc["Video", "result"] == rc.FAIL_VARIATION and not checks.loc["Video", "qualifies"]
    assert not result.available and "at least 2" in result.reason           # 1 qualifying channel


def test_no_channel_or_conversions_explains():
    df = pd.DataFrame({"date": pd.date_range("2025-01-06", periods=10), "spend": 1.0})
    result = rc.response_analysis(df)
    assert not result.available and result.reason == rc.REASON_NO_DATA


def test_incident_weeks_are_left_out():
    spend, conv = power_data(b=0.6, weeks=30, noise=0.01)
    conv = conv.copy()
    conv[5] = conv[5] * 20                                     # a planted outlier week
    df = daily_rows({"Paid Search": (spend, conv), "Video": power_data(b=0.5, weeks=30, seed=2)})
    week = pd.Timestamp("2025-01-06") + pd.Timedelta(days=35)
    incidents = pd.DataFrame([{"incident_id": "I01", "entity_type": "campaign", "entity": "Paid Search A",
                               "period_start": week, "period_end": week + pd.Timedelta(days=6)},
                              {"incident_id": "I02", "entity_type": "region", "entity": "East",
                               "period_start": week, "period_end": week}])
    weeks = rc.incident_weeks(df, incidents, "Paid Search")
    assert list(weeks) == [week] and weeks[week] == ["I01"]
    result = rc.response_analysis(df, incidents)
    ps = result.curves["Paid Search"]
    assert ps.excluded_incidents == ["I01"] and ps.weeks_used == 29
    assert ps.b == pytest.approx(0.6, abs=0.03)                # the outlier did not bend the curve


def test_performance_channel_is_never_fitted_and_capped():
    spend, conv = power_data(weeks=30)
    aff = rc.fit_channel(weekly(spend, conv), "Affiliate")
    assert rc.channel_pricing("Affiliate") == "performance" and aff.status == rc.PERFORMANCE
    assert aff.b is None and aff.boot_b is None                               # never curve-fitted
    assert aff.cpa == pytest.approx(spend[-8:].sum() / conv[-8:].sum())
    assert aff.cap_conversions == pytest.approx(1.2 * conv.max())
    huge = aff.max_spend * 100
    assert aff.conversions(huge) == pytest.approx(aff.cap_conversions)
    assert aff.marginal_cost(huge) is None and aff.marginal_cost(aff.base_spend) == pytest.approx(aff.cpa)


# --- Scenarios ------------------------------------------------------------------------------------

@pytest.fixture(scope="module")
def three():
    df = daily_rows({"Paid Search": power_data(b=0.7, a=1.0, weeks=30, seed=3),
                     "Video": power_data(b=0.4, a=5.0, weeks=30, seed=4),
                     "Affiliate": power_data(b=0.9, a=0.02, weeks=30, seed=5)})
    result = rc.response_analysis(df)
    assert result.available and set(result.usable) == {"Paid Search", "Video", "Affiliate"}
    return result


def test_move_budget_conserves_total_and_respects_guardrails(three):
    base = rc.baseline(three)
    spend, warnings = rc.move_budget(three, "Video", "Paid Search", 15)
    assert sum(spend.values()) == pytest.approx(sum(base.values()))
    assert spend["Video"] == pytest.approx(base["Video"] * 0.85) and not warnings
    big, warnings = rc.move_budget(three, "Paid Search", "Video", 50)
    assert sum(big.values()) == pytest.approx(sum(base.values()))
    video = three.curves["Video"]
    assert big["Video"] <= video.guard_high + 1e-6
    if base["Paid Search"] * 0.5 > video.guard_high - base["Video"]:
        assert warnings and "could be moved" in warnings[0]


def test_guardrails_cap_sliders(three):
    spend, warnings = rc.adjust_channels(three, {"Video": 50, "Paid Search": -50})
    for c, s in spend.items():
        cv = three.curves[c]
        assert cv.guard_low - 1e-6 <= s <= cv.guard_high + 1e-6
    capped, warnings = rc.apply_guardrails(three, {"Video": 1e12, "Paid Search": 1.0, "Affiliate":
                                                   three.curves["Affiliate"].base_spend})
    assert capped["Video"] == pytest.approx(three.curves["Video"].guard_high)
    assert capped["Paid Search"] == pytest.approx(three.curves["Paid Search"].guard_low)
    assert len(warnings) == 2 and all("capped" in w for w in warnings)


def test_suggest_allocation_limits_and_gain(three):
    base = rc.baseline(three)
    spend = rc.suggest_allocation(three)
    assert sum(spend.values()) == pytest.approx(sum(base.values()))
    limit = cfg.SCENARIO_SUGGEST_MAX_CHANGE
    for c, s in spend.items():
        cv = three.curves[c]
        assert base[c] * (1 - limit) - 1e-6 <= s <= base[c] * (1 + limit) + 1e-6
        assert cv.guard_low - 1e-6 <= s <= cv.guard_high + 1e-6
    before = rc.evaluate(three, base).totals["conversions_after"]
    after = rc.evaluate(three, spend).totals["conversions_after"]
    assert after >= before
    assert rc.suggest_allocation(three) == spend                  # deterministic


def test_evaluate_totals_and_ranges(three):
    s = rc.evaluate(three, rc.move_budget(three, "Video", "Paid Search", 15)[0], "Test")
    t = s.totals
    assert t["spend_after"] == pytest.approx(t["spend_before"])
    assert t["conversions_low"] <= t["conversions_after"] <= t["conversions_high"]
    assert t["cac_after"] == pytest.approx(t["spend_after"] / t["conversions_after"])
    assert t["roas_after"] == pytest.approx(t["revenue_after"] / t["spend_after"])
    assert not rc.scenario_display(s).empty


# --- Kalpa sample, exports, page -----------------------------------------------------------------

@pytest.fixture(scope="module")
def kalpa(tmp_path_factory):
    raw, report = load_raw(SAMPLE_DIR / "marketing_clean.csv")
    return run_pipeline(prepare(raw, report), db_path=tmp_path_factory.mktemp("db") / "k.db")


def test_kalpa_data_check_and_curves(kalpa):
    a = kalpa.analysis
    assert "response_curves" not in a.tables()                               # evidence pack unchanged
    r = a.response_curves
    print("\n" + rc.checks_table(r).to_string())
    assert r.available and len(r.checks) == 5 and r.checks["qualifies"].all()
    assert "Email" not in set(r.checks["channel"])
    assert r.curves["Affiliate"].status == rc.PERFORMANCE
    assert {c for c, cv in r.curves.items() if cv.status == rc.FITTED} == {
        "Paid Search", "Paid Social", "Video", "Professional Network"}
    assert "I08" in r.curves["Video"].excluded_incidents                      # the planted spend spike


def test_exports_include_budget_efficiency(kalpa, tmp_path):
    a = kalpa.analysis
    path = tmp_path / "w.xlsx"
    assert "Response_Curves" in generate_workbook(a, path)
    values = [c.value for row in load_workbook(path)["Response_Curves"].iter_rows() for c in row if c.value is not None]
    assert "Paid Search" in values and any(isinstance(v, float) for v in values)
    buffer = io.BytesIO()
    generate_pdf(a, buffer)
    text = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(buffer.getvalue())).pages)
    assert "Budget Efficiency" in text and "Scenario" not in text.split("Budget Efficiency")[1][:1500]
    scenario = rc.evaluate(a.response_curves, rc.suggest_allocation(a.response_curves), "Suggested")
    out = io.BytesIO()
    scenario_workbook(a.response_curves, scenario, out)
    ws = load_workbook(io.BytesIO(out.getvalue()))["Scenario"]
    assert ws["A1"].value == "Scenario: Suggested"


def _app(output, tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    monkeypatch.setattr(dbc, "settings", Settings(database_path=str(tmp_path / "app.db")))
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=240).run()
    at.session_state["output"] = output
    return at


def test_page_renders_and_reacts(kalpa, tmp_path, monkeypatch):
    at = _app(kalpa, tmp_path, monkeypatch)
    at.sidebar.radio(key="page").set_value("Scenario Planner").run()
    assert not at.exception
    assert any(m.label == "Conversions per Week" for m in at.metric)
    at.selectbox(key="scenario_from").set_value("Video").run()
    at.selectbox(key="scenario_to").set_value("Paid Search").run()
    assert not at.exception and any("Move 15% of Video to Paid Search" in md.value for md in at.markdown)
    at.radio(key="scenario_mode").set_value("Adjust Each Channel").run()
    at.slider(key="scenario_adj_Video").set_value(50).run()
    assert not at.exception and any("Weekly Total" in md.value for md in at.markdown)
    at.radio(key="scenario_mode").set_value("Suggest Allocation").run()
    at.selectbox(key="scenario_curve_channel").set_value("Affiliate").run()
    assert not at.exception


def test_page_explains_when_the_check_fails(kalpa, tmp_path, monkeypatch):
    monkeypatch.setattr(cfg, "RESPONSE_MIN_CV", 0.9)                          # nothing varies that much
    failed = rc.response_analysis(kalpa.clean.clean_df, kalpa.analysis.incidents)
    monkeypatch.setattr(kalpa.analysis, "response_curves", failed)          # restored after the test
    at = _app(kalpa, tmp_path, monkeypatch)
    at.sidebar.radio(key="page").set_value("Scenario Planner").run()
    assert not at.exception
    assert not any(m.label == "Conversions per Week" for m in at.metric)
    assert any("passed the data check" in md.value for md in at.markdown)


def test_ai_recommendation_prefills_the_planner(kalpa, tmp_path, monkeypatch):
    """The demo seed's budget move (Meta - Digital Marketing Prospecting -> Affiliate, 15%) opens the
    planner pre-filled: FROM = that campaign's channel, TO = the best-ROAS other channel cited."""
    from ai.context_builder import build_evidence
    from dashboard import layout
    monkeypatch.setattr(ai.demo_seed, "SEED_PATH", ROOT / "data" / "demo" / "kalpa_demo_seed.json")
    monkeypatch.setattr(dashboard.demo, "DEMO_SNAPSHOT_PATH", ROOT / "data" / "demo" / "kalpa_demo_snapshot.pkl.gz")
    import json
    seed = json.loads((ROOT / "data" / "demo" / "kalpa_demo_seed.json").read_text(encoding="utf-8"))
    recs = seed["ai"]["insights"]["recommendations"]
    pack = build_evidence(kalpa.analysis, model="test-model")
    channels = list(kalpa.analysis.response_curves.usable)
    moves = [layout.ai_scenario_move(r, pack, channels) for r in recs]
    assert moves[0] is None                                                  # not a budget move
    assert ("Paid Social", "Affiliate", 15.0) in moves
    from streamlit.testing.v1 import AppTest
    monkeypatch.setattr(dbc, "settings", Settings(database_path=str(tmp_path / "app.db")))
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=240).run()
    at.button(key="load_sample").click().run()
    at.sidebar.radio(key="page").set_value("AI Insights").run()
    buttons = [b for b in at.button if b.label == "Open in Scenario Planner"]
    assert buttons, "no Open in Scenario Planner button"
    buttons[0].click().run()
    assert not at.exception
    assert at.session_state["page"] == "Scenario Planner"
    assert at.selectbox(key="scenario_from").value in channels
    assert any("Pre-filled from an AI recommendation" in s.value for s in at.success)
