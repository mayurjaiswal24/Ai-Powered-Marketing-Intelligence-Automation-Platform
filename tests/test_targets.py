"""U5 Targets vs actual: status in both directions at the boundaries, monthly spend, partial months,
save/load/clear per column layout, public-mode isolation, exports, the page and the Executive
Overview line (unchanged without targets)."""

import io
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

import config.settings as config_settings
import database.connection as dbc
from analytics import targets as tg
from analytics.kpis import compute_kpis
from config.settings import SPEND_ON_TARGET_BAND, TARGET_TOLERANCE, Settings
from dashboard.pipeline import load_raw, prepare, run_pipeline
from database import repository
from reports.excel_report import generate_workbook
from reports.pdf_report import generate_pdf

ROOT = Path(__file__).resolve().parent.parent
SAMPLE_DIR = ROOT / "data" / "sample"


# --- Status -------------------------------------------------------------------------------------

@pytest.mark.parametrize("actual,expected", [
    (90.0, tg.ON), (100.0, tg.ON),              # at or below target
    (110.0, tg.WITHIN),                         # exactly 10% above: still within
    (110.01, tg.OFF), (150.0, tg.OFF),
])
def test_lower_is_better_boundaries(actual, expected):
    assert tg.target_status(actual, 100.0, "cpl") == expected
    assert tg.target_status(actual, 100.0, "cac") == expected


@pytest.mark.parametrize("actual,expected", [
    (5.5, tg.ON), (5.0, tg.ON),                 # at or above target
    (4.5, tg.WITHIN),                           # exactly 10% below: still within
    (4.4999, tg.OFF), (1.0, tg.OFF),
])
def test_higher_is_better_boundaries(actual, expected):
    assert tg.target_status(actual, 5.0, "roas") == expected


@pytest.mark.parametrize("actual,expected", [
    (100.0, tg.ON), (105.0, tg.ON), (95.0, tg.ON),          # within +/-5%
    (105.01, tg.WITHIN), (110.0, tg.WITHIN), (90.0, tg.WITHIN),
    (110.01, tg.OVER), (89.99, tg.UNDER),
])
def test_monthly_spend_band(actual, expected):
    assert tg.target_status(actual, 100.0, "monthly_spend") == expected


def test_settings_and_missing_values():
    assert TARGET_TOLERANCE == 0.10 and SPEND_ON_TARGET_BAND == 0.05
    assert tg.target_status(None, 100.0, "cpl") == tg.NOT_AVAILABLE
    assert tg.target_status(float("nan"), 100.0, "roas") == tg.NOT_AVAILABLE
    # Monthly leads/conversions/revenue follow the registry: higher is better.
    assert tg.target_status(95.0, 100.0, "monthly_leads") == tg.WITHIN
    assert tg.target_status(120.0, 100.0, "monthly_revenue") == tg.ON
    # A headroom target of 0 ("at least break-even").
    assert tg.target_status(-3.0, 0.0, "roas_headroom") == tg.OFF
    assert tg.target_status(3.0, 0.0, "roas_headroom") == tg.ON
    assert tg.clean_targets({"cpl": None, "roas": 4, "unknown": 1}) == {"roas": 4.0}


# --- Analysis on a small hand-made dataset ----------------------------------------------------------

def fixture_df() -> pd.DataFrame:
    """January complete (31 days), February partial (3 days). Spend 100/day, 2 leads/day."""
    dates = list(pd.date_range("2026-01-01", "2026-01-31")) + list(pd.date_range("2026-02-01", "2026-02-03"))
    rows = [{"date": d, "month": d.strftime("%Y-%m"), "campaign_name": "A",
             "channel": "Paid Search" if i % 2 else "Email", "spend": 100.0, "leads": 2.0,
             "clicks": 20.0, "impressions": 1000.0} for i, d in enumerate(dates)]
    return pd.DataFrame(rows)


def test_targets_analysis_ratios_months_and_channels():
    df = fixture_df()
    result = tg.targets_analysis(df, {"cpl": 50.0, "monthly_spend": 3000.0, "ctr": None})
    sc = result.scorecard.set_index("metric")
    assert list(sc.index) == ["cpl", "monthly_spend"]           # blank = no target, nothing shown
    assert sc.loc["cpl", "actual"] == pytest.approx(compute_kpis(df)["cpl"].value) == pytest.approx(50.0)
    assert sc.loc["cpl", "status"] == tg.ON
    # January is the latest full month: 31 x 100 = 3,100 vs 3,000 = +3.3% -> On Target.
    assert sc.loc["monthly_spend", "actual"] == pytest.approx(3100.0)
    assert sc.loc["monthly_spend", "basis"] == "Jan 2026 (latest full month)"
    assert sc.loc["monthly_spend", "status"] == tg.ON
    monthly = result.monthly.set_index("period")
    assert monthly.loc["2026-02", "status"] == tg.PARTIAL        # 3 days are not "under target"
    channels = result.channels.set_index("channel")
    assert set(channels.index) == {"Paid Search", "Email (owned)"}
    assert tg.summary_lines(result)[0] == "2 of 2 targets are on target."


def test_no_targets_means_nothing_shown():
    assert tg.targets_analysis(fixture_df(), {}) is None
    assert tg.targets_analysis(fixture_df(), {"cpl": None}) is None
    assert tg.overview_notes(compute_kpis(fixture_df()), {}, ["cpl"]) == {}


def test_overview_note_text():
    kpis = compute_kpis(fixture_df())
    assert tg.overview_notes(kpis, {"cpl": 9000.0}, ["spend", "cpl"]) == {"cpl": "Target ₹9,000 · On Target"}


def test_available_metrics_follow_the_data():
    df = fixture_df()
    keys = tg.available_metrics(df)
    assert "cpl" in keys and "ctr" in keys and "monthly_spend" in keys and "monthly_leads" in keys
    assert "roas" not in keys and "roas_headroom" not in keys and "monthly_revenue" not in keys


# --- Saving and loading --------------------------------------------------------------------------------

def test_save_load_clear_and_layout_matching(tmp_path):
    db = tmp_path / "t.db"
    state = {}
    tg.save_targets(state, "layout-A", {"cpl": 250.0, "roas": None}, public=False, db_path=db)
    assert repository.load_targets("layout-A", db_path=db) == {"cpl": 250.0}
    # A new session with the same column layout gets them back; another layout gets none.
    assert tg.load_targets({}, "layout-A", public=False, db_path=db) == {"cpl": 250.0}
    assert tg.load_targets({}, "layout-B", public=False, db_path=db) == {}
    # Saving again replaces (a metric left blank loses its target).
    tg.save_targets(state, "layout-A", {"roas": 5.0}, public=False, db_path=db)
    assert repository.load_targets("layout-A", db_path=db) == {"roas": 5.0}
    tg.clear_targets(state, "layout-A", public=False, db_path=db)
    assert repository.load_targets("layout-A", db_path=db) == {}
    assert tg.load_targets(state, "layout-A", public=False, db_path=db) == {}


def test_public_mode_targets_stay_in_the_session(tmp_path):
    db = tmp_path / "p.db"
    visitor_1 = {}
    tg.save_targets(visitor_1, "layout-A", {"cpl": 250.0}, public=True, db_path=db)
    assert tg.load_targets(visitor_1, "layout-A", public=True, db_path=db) == {"cpl": 250.0}
    assert repository.load_targets("layout-A", db_path=db) == {}             # nothing written
    assert tg.load_targets({}, "layout-A", public=True, db_path=db) == {}   # visitor 2 sees none


# --- The sample dataset: layout key, exports, page, overview ---------------------------------------

KALPA_TARGETS = {"cpl": 240.0, "cac": 9200.0, "roas": 5.0, "ctr": 0.78, "lead_to_conversion_rate": 2.9,
                 "monthly_spend": 6200000.0, "monthly_leads": 25000.0}


@pytest.fixture(scope="module")
def kalpa(tmp_path_factory):
    raw, report = load_raw(SAMPLE_DIR / "marketing_clean.csv")
    return run_pipeline(prepare(raw, report), db_path=tmp_path_factory.mktemp("db") / "k.db")


def test_kalpa_layout_key_matches_mapping_cache(kalpa):
    from ai.mapping import layout_key
    raw, _ = load_raw(SAMPLE_DIR / "marketing_clean.csv")
    assert kalpa.layout_key == layout_key(raw.columns)
    result = tg.targets_analysis(kalpa.clean.clean_df, KALPA_TARGETS, kalpa.analysis.capabilities,
                                 kalpa.analysis.profitability, kpis=kalpa.analysis.kpis)
    sc = result.scorecard.set_index("metric")
    assert sc.loc["roas", "actual"] == pytest.approx(kalpa.analysis.kpis["roas"].value)
    assert sc.loc["cpl", "status"] == tg.OFF                   # ₹268 vs ₹240 is 11.8% over


def test_exports_only_with_targets(kalpa, tmp_path):
    a = kalpa.analysis
    result = tg.targets_analysis(kalpa.clean.clean_df, KALPA_TARGETS, a.capabilities, a.profitability,
                                 kpis=a.kpis)
    path = tmp_path / "t.xlsx"
    assert "Targets" in generate_workbook(a, path, targets=result)
    ws = load_workbook(path)["Targets"]
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "CPL" in values and 240.0 in values                  # real numbers, not text
    assert "Targets" not in generate_workbook(a, tmp_path / "n.xlsx")
    buffer = io.BytesIO()
    generate_pdf(a, buffer, targets=result)
    text = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(buffer.getvalue())).pages)
    assert "Performance Vs Targets" in text or "Performance vs Targets" in text
    plain = io.BytesIO()
    generate_pdf(a, plain)
    assert "Targets" not in " ".join(p.extract_text() for p in PdfReader(io.BytesIO(plain.getvalue())).pages)


def _app(kalpa, tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    monkeypatch.setattr(dbc, "settings", Settings(database_path=str(tmp_path / "app.db")))
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=240).run()
    at.session_state["output"] = kalpa
    return at


def _target_captions(at) -> list[str]:
    return [c.value for c in at.caption if c.value.startswith("Target ")]


def test_overview_unchanged_without_targets(kalpa, tmp_path, monkeypatch):
    at = _app(kalpa, tmp_path, monkeypatch)
    at.sidebar.radio(key="page").set_value("Executive Overview").run()
    assert not at.exception and _target_captions(at) == []


def test_targets_page_save_and_overview_line(kalpa, tmp_path, monkeypatch):
    at = _app(kalpa, tmp_path, monkeypatch)
    at.sidebar.radio(key="page").set_value("Targets").run()
    assert not at.exception
    key = f"target_{kalpa.layout_key[-12:]}_cpl"
    assert at.number_input(key=key).value is None                # blank by default
    at.number_input(key=key).set_value(9000.0)
    at.button(key="FormSubmitter:" + f"targets_form_{kalpa.layout_key[-12:]}-Save Targets").click().run()
    assert not at.exception
    assert any("CPL" in str(df.value) for df in at.dataframe)
    assert repository.load_targets(kalpa.layout_key, db_path=tmp_path / "app.db") == {"cpl": 9000.0}
    at.sidebar.radio(key="page").set_value("Executive Overview").run()
    assert _target_captions(at) == ["Target ₹9,000 · On Target"]


def test_targets_page_public_mode_does_not_save(kalpa, tmp_path, monkeypatch):
    import dataclasses
    monkeypatch.setattr(config_settings, "settings",
                        dataclasses.replace(config_settings.settings, public_mode=True))
    at = _app(kalpa, tmp_path, monkeypatch)
    at.sidebar.radio(key="page").set_value("Targets").run()
    assert not at.exception
    assert any("this browser session only" in c.value for c in at.caption)
