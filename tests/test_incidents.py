"""Review changes (before Phase 13): owned channels, incidents, trend notes, attribution note,
prompt v2."""

import json
from pathlib import Path

import pandas as pd
import pytest

from ai.prompts import PROMPT_VERSION, SYSTEM_INSTRUCTIONS
from analytics.incidents import build_incidents, match_planted
from config.fields import channel_type
from dashboard.pipeline import SAMPLE_DIR, load_raw, prepare, run_pipeline
from ingestion.validator import ATTRIBUTION_NOTE

ROOT = Path(__file__).resolve().parent.parent


@pytest.fixture(scope="module")
def clean(tmp_path_factory):
    raw, report = load_raw(SAMPLE_DIR / "marketing_clean.csv")
    return run_pipeline(prepare(raw, report), db_path=tmp_path_factory.mktemp("i") / "i.db")


@pytest.fixture(scope="module")
def truth():
    return json.loads((SAMPLE_DIR / "ground_truth_anomalies.json").read_text(encoding="utf-8"))["anomalies"]


# --- Data: realistic email cost ----------------------------------------------------------------

def test_email_roas_is_realistic(clean):
    email = clean.analysis.channels.set_index("channel").loc["Email"]
    assert 15 <= email["roas"] <= 30
    assert email["spend"] > 5_00_000                     # platform + content cost, not near zero


# --- Owned channels ----------------------------------------------------------------------------

def test_channel_type_config():
    assert channel_type("Email") == "owned" and channel_type(" email ") == "owned"
    assert channel_type("Email (in-house)") == "owned"
    assert channel_type("Paid Search") == "paid" and channel_type("Something new") == "paid"


def test_owned_channels_excluded_from_rankings_and_indices(clean):
    ch = clean.analysis.channels.set_index("channel")
    assert ch.loc["Email", "channel_type"] == "owned"
    assert pd.isna(ch.loc["Email", "roas_index"]) and pd.isna(ch.loc["Email", "cpl_index"])
    assert ch.loc["Paid Search", "roas_index"] > 0
    findings = {f.id: f for f in clean.analysis.findings}
    assert findings["F-best-channel"].entity != "Email" and findings["F-weakest-channel"].entity != "Email"
    assert "paid channels" in findings["F-best-channel"].text
    owned = findings["F-owned-Email"]
    assert "owned channel" in owned.text and "not ranked" in owned.text
    assert findings["F-weakest-campaign"].entity.split(" - ")[0] != "Email"


# --- Incidents ---------------------------------------------------------------------------------

def _planted_targets(g):
    """(entity_type, entity) pairs a planted anomaly can be matched on."""
    if g["campaign_names"] != "all":
        return [("campaign", n) for n in g["campaign_names"]]
    if g["region"] != "all":
        return [("region", g["region"])]
    platforms = g["platform"] if isinstance(g["platform"], list) else [g["platform"]]
    return [("platform", p) for p in platforms]


def test_incident_recall_is_complete(clean, truth):
    incidents = clean.analysis.incidents
    found = 0
    lines = []
    for g in truth:
        hits = [r for _, r in incidents.iterrows()
                if any(match_planted(r, et, e, g["start_date"], g["end_date"]) for et, e in _planted_targets(g))]
        found += bool(hits)
        lines.append(f"  {g['id']} {'DETECTED' if hits else 'MISSED  '} "
                     + ", ".join(f"{h['incident_id']} {h['entity']}" for h in hits[:2]))
    print(f"\nIncidents: {len(incidents)} (from {len(clean.analysis.anomalies)} flags); "
          f"planted-anomaly recall at incident level: {found}/{len(truth)}")
    print("\n".join(lines))
    assert found == len(truth)
    assert len(incidents) < len(clean.analysis.anomalies) / 2


def test_outage_absorbs_its_campaign_flags(clean):
    inc = clean.analysis.incidents
    outage = inc[(inc["entity_type"] == "platform") & (inc["entity"] == "Google Ads")].iloc[0]
    assert outage["n_related"] >= 3
    related = {(f["entity_type"], f["entity"]) for f in outage["related"]}
    assert all(t in ("campaign", "channel") for t, _ in related)
    assert outage["impact_label"] == "spend with no tracked results" and outage["impact_inr"] > 0
    # Google Ads campaign flags in that week are no longer separate top-level incidents.
    google = set(clean.clean.clean_df.loc[clean.clean.clean_df["platform"] == "Google Ads", "campaign_name"])
    sep = inc[(inc["entity_type"] == "campaign") & inc["entity"].isin(google)
              & (inc["period_start"] <= pd.Timestamp("2026-09-09")) & (inc["period_end"] >= pd.Timestamp("2026-09-08"))]
    assert sep.empty, sep[["entity", "metrics"]]


def test_lower_cost_is_not_an_improvement_when_conversions_collapse(clean):
    inc = clean.analysis.incidents
    for _, r in inc.iterrows():
        flags = r["flags"] + r["related"]
        collapsed = any(f["metric"] in ("lead_to_conversion_rate", "revenue") and (f["change_pct"] or 0) < 0
                        for f in flags) or any(f["method"] == "tracking_outage_rule" for f in flags)
        if collapsed:
            for f in flags:
                if f["metric"] in ("cpc", "cpl") and (f["change_pct"] or 0) < 0:
                    assert f["sentiment"] == "check" and "not an improvement" in f["note"]


def test_incidents_ranked_by_impact_and_same_entity_grouped(clean):
    inc = clean.analysis.incidents
    assert inc["impact_inr"].is_monotonic_decreasing
    assert inc["rank"].tolist() == list(range(1, len(inc) + 1))
    east = inc[(inc["entity_type"] == "region") & (inc["entity"] == "East")].iloc[0]
    assert {"Lead-to-conversion rate", "Revenue"} <= set(east["metrics"].split(", "))
    assert east["sentiment"] == "negative" and east["impact_direction"] == "loss"


def test_build_incidents_empty():
    assert build_incidents(pd.DataFrame(), pd.DataFrame()).empty


# --- Trend note, attribution note -------------------------------------------------------------

def test_mover_findings_mention_incidents(clean):
    movers = [f for f in clean.analysis.findings if f.id in ("F-riser", "F-faller")]
    assert movers
    noted = [f for f in movers if "contains a detected incident" in f.text]
    assert noted, [f.text for f in movers]
    assert any("Google Ads" in f.text for f in noted)          # the 8-9 Sep tracking outage


def test_attribution_limitation_everywhere(clean, tmp_path):
    from pypdf import PdfReader
    from reports.pdf_report import generate_pdf
    assert ATTRIBUTION_NOTE in clean.analysis.notes
    path = tmp_path / "a.pdf"
    generate_pdf(clean.analysis, path)
    text = " ".join("".join(p.extract_text() for p in PdfReader(str(path)).pages).split())
    assert "upper-funnel channels" in text and "Owned channels (not ranked)" in text
    no_rev_raw, rep = load_raw(SAMPLE_DIR / "marketing_no_revenue.csv")
    no_rev = run_pipeline(prepare(no_rev_raw, rep), db_path=tmp_path / "n.db").analysis
    assert ATTRIBUTION_NOTE not in no_rev.notes                 # no ROAS, no ROAS caveat


def test_excel_has_incidents_and_owned_block(clean, tmp_path):
    from openpyxl import load_workbook
    from reports.excel_report import generate_workbook
    path = tmp_path / "w.xlsx"
    sheets = generate_workbook(clean.analysis, path)
    assert sheets.index("Incidents") < sheets.index("Anomalies")
    wb = load_workbook(path)
    values = [c.value for row in wb["Channel_Analysis"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any(v.startswith("Owned channels") for v in values)
    inc = wb["Incidents"]
    assert isinstance(inc.cell(4, [c.value for c in inc[3]].index("Estimated impact (₹)") + 1).value, float)


# --- Prompt v2 ---------------------------------------------------------------------------------

def test_prompt_keeps_v2_rules():
    assert PROMPT_VERSION >= "v2"            # the v2 rules stay in later versions
    text = " ".join(SYSTEM_INSTRUCTIONS.lower().split())
    for rule in ("connect at", "never recommend increasing", "click-to-lead rate points to landing page",
                 "not the landing page", "attribution limitation", "from and to", "not a customer segment",
                 "incident"):
        assert rule in text, rule
