"""U6 Budget pacing and forecast: pacing maths (incl. a mid-month as-of date), user budget, no
budget, status band, forecast on a known trend and a flat series, insufficient history, zeros,
horizon limit, determinism, range coverage on the Kalpa sample, exports and the page. The AI
evidence pack is unchanged (tests/test_demo.py pins its fingerprint)."""

import io
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

import database.connection as dbc
from analytics import forecast as fc
from analytics import pacing as pc
from analytics.kpis import pacing_metrics, pacing_status
from config.settings import FORECAST_MIN_WEEKS, PACING_BAND, Settings
from dashboard.pipeline import load_raw, prepare, run_pipeline
from reports.excel_report import generate_workbook
from reports.pdf_report import generate_pdf

ROOT = Path(__file__).resolve().parent.parent
SAMPLE_DIR = ROOT / "data" / "sample"


def january(spend_by_day, budget: float | None = 100.0, days: int = 31) -> pd.DataFrame:
    """One row per day in January 2026 (two channels alternate); planned budget 100 per day."""
    rows = []
    for i, d in enumerate(pd.date_range("2026-01-01", periods=days)):
        row = {"date": d, "month": "2026-01", "campaign": "A" if i % 2 else "B",
               "channel": "Paid Search" if i % 2 else "Email", "spend": float(spend_by_day(d.day))}
        if budget is not None:
            row["budget"] = budget
        rows.append(row)
    return pd.DataFrame(rows)


# --- Pacing maths -----------------------------------------------------------------------------------

@pytest.mark.parametrize("ratio,expected", [
    (1.0, "On Pace"),
    (1.05, "On Pace"), (0.95, "On Pace"),               # the band boundaries count as on pace
    (1.0501, "Overspending"), (0.9499, "Underspending"), (None, "Not available"),
])
def test_status_band(ratio, expected):
    assert PACING_BAND == 0.05
    assert pacing_status(ratio) == expected


def test_full_month_utilisation_and_variance():
    result = pc.pacing_analysis(january(lambda day: 90.0))
    assert result.available and result.source == pc.SOURCE_COLUMN and not result.assumed
    m = result.monthly.iloc[0]
    assert m["budget"] == pytest.approx(3100) and m["spend"] == pytest.approx(2790)
    assert m["budget_utilisation"] == pytest.approx(90.0) and m["variance"] == pytest.approx(-310)
    assert m["status"] == "Underspending"
    p = result.overall                                   # as of the last date: nothing left to plan
    assert p["remaining_planned"] == 0 and p["projected_spend"] == pytest.approx(2790)
    assert p["projected_variance"] == pytest.approx(-310)


def test_mid_month_as_of():
    # Days 1-10 spend 100, days 11-31 spend 120; as of 15 Jan.
    df = january(lambda day: 100.0 if day <= 10 else 120.0)
    p = pc.pacing_analysis(df, as_of="2026-01-15").overall
    assert p["planned_to_date"] == pytest.approx(1500)
    assert p["spend_to_date"] == pytest.approx(1600)
    assert p["pacing_ratio"] == pytest.approx(1600 / 1500) and p["status"] == "Overspending"
    recent = (2 * 100 + 5 * 120) / 700                    # 9-15 Jan
    assert p["recent_utilisation"] == pytest.approx(recent)
    assert p["remaining_planned"] == pytest.approx(1600) and p["month_budget"] == pytest.approx(3100)
    assert p["projected_spend"] == pytest.approx(1600 + 1600 * recent)
    assert p["projected_variance"] == pytest.approx(1600 + 1600 * recent - 3100)
    assert p["day"] == 15 and p["days_in_month"] == 31 and p["estimated_days"] == 0


def test_data_ending_mid_month_plans_the_rest_at_the_recent_average():
    p = pc.pacing_analysis(january(lambda day: 100.0, days=15)).overall
    assert p["estimated_days"] == 16 and p["remaining_planned"] == pytest.approx(1600)
    assert p["status"] == "On Pace" and p["projected_spend"] == pytest.approx(3100)
    monthly = pc.pacing_analysis(january(lambda day: 100.0, days=15)).monthly
    assert monthly.iloc[0]["status"] == pc.PARTIAL        # half a month is not rated


def test_formula_edge_cases():
    none = pacing_metrics(0, 0, 0, 0, 0)
    assert none["pacing_ratio"] is None and none["projected_variance"] is None
    fallback = pacing_metrics(500, 1000, 1000, 0, 0)       # no recent plan -> month-to-date pace
    assert fallback["recent_utilisation"] == pytest.approx(0.5) and fallback["projected_spend"] == pytest.approx(1000)


def test_user_budget_and_labels():
    df = january(lambda day: 100.0, budget=None)
    result = pc.pacing_analysis(df, monthly_budget=3100.0)
    assert result.available and result.assumed and result.source == pc.SOURCE_USER
    assert result.label == "Based on your monthly budget of ₹3,100"
    assert result.overall["planned_to_date"] == pytest.approx(3100) and result.overall["status"] == "On Pace"
    assert result.channels.empty and result.campaigns.empty          # needs a Budget column
    assert any("Based on your monthly budget" in line for line in pc.summary_lines(result))
    # A Budget column in the data always wins over a typed budget.
    assert not pc.pacing_analysis(january(lambda day: 100.0), monthly_budget=9999.0).assumed


def test_no_budget_and_no_dates_explain_why():
    df = january(lambda day: 100.0, budget=None)
    assert pc.pacing_analysis(df).reason == pc.REASON_NO_BUDGET
    assert pc.pacing_analysis(df, monthly_budget=0).reason == pc.REASON_NO_BUDGET
    assert pc.pacing_analysis(df.drop(columns=["date"])).reason == pc.REASON_NO_DATES


def test_by_channel_marks_owned():
    result = pc.pacing_analysis(january(lambda day: 100.0), as_of="2026-01-20")
    assert set(result.channels["channel"]) == {"Paid Search", "Email (owned)"}
    assert result.channels["spend_to_date"].sum() == pytest.approx(result.overall["spend_to_date"])
    assert len(result.campaigns) == 2


# --- Forecast ----------------------------------------------------------------------------------------

def weekly_df(values) -> pd.DataFrame:
    """Daily rows whose weekly totals equal `values` (weeks start on Monday 5 Jan 2026)."""
    rows = []
    for w, total in enumerate(values):
        for d in range(7):
            day = pd.Timestamp("2026-01-05") + pd.Timedelta(days=7 * w + d)
            rows.append({"date": day, "channel": "Paid Search", "spend": total / 7, "leads": total / 70,
                         "conversions": 0.0, "revenue": total * 4 / 7})
    return pd.DataFrame(rows)


def test_known_trend_is_followed():
    values = [1000 + 50 * w for w in range(40)]
    result = fc.forecast_analysis(weekly_df(values))
    assert result.available and result.weeks == 40 and result.horizon == 8
    mf = result.metrics["spend"]
    assert mf.method in ("ma4_trend", "holt") and mf.wape < 1.0
    truth = [1000 + 50 * (40 + h) for h in range(8)]
    assert np.allclose(mf.forecast["forecast"], truth, rtol=0.01)
    assert (mf.forecast["low"] <= mf.forecast["forecast"]).all() and (mf.forecast["forecast"] <= mf.forecast["high"]).all()


def test_flat_series_gives_flat_forecast_and_no_range():
    mf = fc.forecast_analysis(weekly_df([700.0] * 30)).metrics["spend"]
    assert mf.method == "last_value" and mf.wape == pytest.approx(0.0)        # ties -> simplest method
    assert np.allclose(mf.forecast["forecast"], 700.0)
    assert np.allclose(mf.forecast["low"], 700.0) and np.allclose(mf.forecast["high"], 700.0)


def test_insufficient_history():
    result = fc.forecast_analysis(weekly_df([100.0] * (FORECAST_MIN_WEEKS - 1)))
    assert not result.available and result.weeks == FORECAST_MIN_WEEKS - 1
    assert "at least 26 full weeks" in result.reason and "has 25" in result.reason
    assert not fc.forecast_analysis(pd.DataFrame({"spend": [1.0]})).available


def test_zeros_do_not_break_the_forecast():
    result = fc.forecast_analysis(weekly_df([0.0] * 20 + [500.0] * 10))
    zero = result.metrics["conversions"]                       # always zero
    assert zero.wape is None and np.allclose(zero.forecast["forecast"], 0.0)
    assert "no activity" in fc.accuracy_sentence(zero)
    assert (result.metrics["spend"].forecast["low"] >= 0).all()


def test_partial_edge_weeks_are_left_out():
    df = weekly_df([100.0] * 30)
    extra = df.iloc[-1:].assign(date=df["date"].max() + pd.Timedelta(days=1))   # one day of a new week
    assert len(fc.weekly_history(pd.concat([df, extra]))) == 30


def test_horizon_is_capped_and_selectable():
    df = weekly_df([100.0 + w for w in range(30)])
    assert fc.forecast_analysis(df, 12).horizon == 8
    assert len(fc.forecast_analysis(df, 12).metrics["spend"].forecast) == 8
    assert len(fc.forecast_analysis(df, 4).metrics["revenue"].forecast) == 4


def test_forecast_is_deterministic():
    rng = np.random.default_rng(7)
    df = weekly_df(list(1000 + rng.normal(0, 80, 40)))
    a, b = fc.forecast_analysis(df), fc.forecast_analysis(df)
    for m in a.metrics:
        pd.testing.assert_frame_equal(a.metrics[m].forecast, b.metrics[m].forecast)
        assert a.metrics[m].method == b.metrics[m].method


def test_channel_forecast_uses_only_that_channel():
    df = pd.concat([weekly_df([100.0] * 30), weekly_df([900.0] * 30).assign(channel="Email")])
    assert np.allclose(fc.forecast_analysis(df, channel="Email").metrics["spend"].forecast["forecast"], 900.0)
    assert np.allclose(fc.forecast_analysis(df).metrics["spend"].forecast["forecast"], 1000.0)


# --- Kalpa sample ------------------------------------------------------------------------------------

@pytest.fixture(scope="module")
def kalpa(tmp_path_factory):
    raw, report = load_raw(SAMPLE_DIR / "marketing_clean.csv")
    return run_pipeline(prepare(raw, report), db_path=tmp_path_factory.mktemp("db") / "k.db")


def test_kalpa_pacing_and_forecast_in_the_result(kalpa):
    a = kalpa.analysis
    assert "pacing" not in a.tables() and "forecast" not in a.tables()       # evidence pack unchanged
    p = a.pacing
    assert p.available and p.source == pc.SOURCE_COLUMN and len(p.monthly) == 12
    assert p.monthly["budget_utilisation"].between(79, 84).all()             # the sample spends ~81%
    assert p.overall["status"] == "Underspending"
    assert "Email (owned)" in set(p.channels["channel"])
    assert p.channels["spend_to_date"].sum() == pytest.approx(p.overall["spend_to_date"])
    f = a.forecast
    assert f.available and f.weeks >= FORECAST_MIN_WEEKS and set(f.metrics) == set(fc.FORECAST_METRICS)
    mid = pc.pacing_analysis(kalpa.clean.clean_df, as_of="2026-09-15").overall
    assert mid["day"] == 15 and mid["remaining_planned"] > 0 and mid["projected_spend"] > mid["spend_to_date"]


def test_kalpa_range_coverage(kalpa):
    """Replay the last 26 weeks: the share of real weeks inside the shaded range (expected ~70-90%)."""
    coverage = fc.range_coverage(kalpa.clean.clean_df, horizon=8)
    print(f"Kalpa 8-week range coverage: {coverage:.1f}%")
    assert 60.0 <= coverage <= 95.0


def test_exports_include_pacing_and_forecast(kalpa, tmp_path):
    a = kalpa.analysis
    path = tmp_path / "w.xlsx"
    sheets = generate_workbook(a, path)
    assert "Pacing" in sheets and "Forecast" in sheets
    book = load_workbook(path)
    values = [c.value for row in book["Forecast"].iter_rows() for c in row if c.value is not None]
    assert any(isinstance(v, float) for v in values) and "Spend" in values        # real numbers
    buffer = io.BytesIO()
    generate_pdf(a, buffer)
    text = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(buffer.getvalue())).pages)
    assert "Budget Pacing and Outlook" in text or "Budget Pacing And Outlook" in text


def test_user_budget_exports_are_labelled(tmp_path):
    df = january(lambda day: 100.0, budget=None)
    user = pc.pacing_analysis(df, monthly_budget=3100.0)
    raw, report = load_raw(SAMPLE_DIR / "marketing_no_margin.csv")
    a = run_pipeline(prepare(raw, report), db_path=tmp_path / "n.db").analysis
    path = tmp_path / "u.xlsx"
    generate_workbook(a, path, pacing=user)
    first = load_workbook(path)["Pacing"]["A1"].value
    assert first.startswith("Based on your monthly budget of ₹3,100")


# --- Page ----------------------------------------------------------------------------------------------

def _app(kalpa, tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    monkeypatch.setattr(dbc, "settings", Settings(database_path=str(tmp_path / "app.db")))
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=240).run()
    at.session_state["output"] = kalpa
    return at


def test_page_renders_and_reacts(kalpa, tmp_path, monkeypatch):
    at = _app(kalpa, tmp_path, monkeypatch)
    at.sidebar.radio(key="page").set_value("Pacing & Forecast").run()
    assert not at.exception
    assert any(m.label == "Pacing" for m in at.metric)
    at.date_input(key="pacing_as_of").set_value(pd.Timestamp("2026-09-15").date()).run()
    assert not at.exception
    assert any("day 15 of 30" in md.value for md in at.markdown)
    at.selectbox(key="forecast_channel").set_value("Email").run()
    at.selectbox(key="forecast_horizon").set_value(4).run()
    assert not at.exception
    assert any("most accurate" in i.value for i in at.info)
