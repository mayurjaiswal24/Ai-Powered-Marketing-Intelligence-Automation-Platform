"""Phase 10 tests: the Excel workbook is generated, adaptive, and holds real, matching numbers."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

import database.connection as dbc
import reports.excel_report as excel_report
from config.settings import Settings
from dashboard.pipeline import SAMPLE_DIR, load_raw, prepare, run_pipeline
from database import repository as repo
from reports.excel_report import export_workbook, generate_workbook

ROOT = Path(__file__).resolve().parent.parent
DATASETS = {
    "clean": ("marketing_clean.csv", None),
    "messy": ("marketing_messy.xlsx", None),
    "no_revenue": ("marketing_no_revenue.csv", None),
    "no_margin": ("marketing_no_margin.csv", None),
    "meta": ("meta_ads_export_style.csv", {"Results": "leads"}),
}
ALWAYS = ["Executive_KPIs", "Clean_Data", "KPI_Analysis", "Data_Quality", "Methodology"]


@pytest.fixture(scope="module")
def books(tmp_path_factory):
    folder = tmp_path_factory.mktemp("xlsx")
    out = {}
    for name, (file, overrides) in DATASETS.items():
        raw, report = load_raw(SAMPLE_DIR / file)
        output = run_pipeline(prepare(raw, report, overrides), db_path=folder / "t.db")
        path = folder / f"{name}.xlsx"
        sheets = generate_workbook(output.analysis, path, output.clean.clean_df,
                                   output.clean.quality_log_df)
        out[name] = (output, sheets, load_workbook(path))
    return out


def kpi_rows(wb) -> dict[str, object]:
    """KPI label -> cell (from the Executive_KPIs table)."""
    ws = wb["Executive_KPIs"]
    rows, started = {}, False
    for row in ws.iter_rows(min_row=1, max_col=2):
        label, value = row[0].value, row[1]
        if label == "KPI":
            started = True
            continue
        if started and label:
            if label == "Not available for this dataset":
                break
            rows[label] = value
    return rows


def test_workbook_generates_for_every_dataset(books):
    for name, (_, sheets, wb) in books.items():
        assert wb.sheetnames == sheets, name
        for sheet in ALWAYS:
            assert sheet in sheets, (name, sheet)


def test_sheet_list_follows_capabilities(books):
    clean = books["clean"][1]
    for sheet in ["Campaign_Analysis", "Channel_Analysis", "Segment_Analysis", "Funnel_Analysis",
                  "Trend_Analysis", "Anomalies"]:
        assert sheet in clean
    assert "AI_Insights" not in clean                        # no AI output yet
    meta = books["meta"][1]
    assert "Channel_Analysis" not in meta and "Segment_Analysis" not in meta
    assert "Campaign_Analysis" in meta and "Funnel_Analysis" in meta


def test_key_kpi_cells_are_numbers_equal_to_analysis(books):
    output, _, wb = books["clean"]
    cells = kpi_rows(wb)
    for k in output.analysis.kpis.values():
        if not k.available:
            assert k.label not in cells
            continue
        cell = cells[k.label]
        assert isinstance(cell.value, (int, float)), k.key          # a number, not text
        assert cell.value == pytest.approx(k.value, rel=1e-12), k.key
    assert "₹" in cells["Spend"].number_format
    assert cells["ROAS"].number_format == '0.00"x"'
    assert cells["CTR"].number_format == '0.0"%"'


def test_unavailable_kpis_listed_not_zero(books):
    output, _, wb = books["no_revenue"]
    cells = kpi_rows(wb)
    assert "Revenue" not in cells and "ROAS" not in cells and "ROI" not in cells
    text = [c.value for row in wb["Executive_KPIs"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any("no revenue field was detected" in t for t in text)
    assert "Revenue" not in [c.value for c in wb["Channel_Analysis"][3]]


def test_tables_have_real_numbers_filters_and_frozen_headers(books):
    output, _, wb = books["clean"]
    ws = wb["Channel_Analysis"]
    headers = [c.value for c in ws[3]]
    spend_col = headers.index("Spend") + 1
    channels = output.analysis.channels.reset_index(drop=True)
    for i, expected in enumerate(channels["spend"], start=4):
        assert ws.cell(i, spend_col).value == pytest.approx(expected)
    assert ws.auto_filter.ref and ws.freeze_panes == "B4"
    clean = wb["Clean_Data"]
    assert clean.max_row == len(output.clean.clean_df) + 1
    assert clean["A1"].font.bold
    date_col = [c.value for c in clean[1]].index("date") + 1
    assert clean.cell(2, date_col).is_date


def test_messy_quality_log_included(books):
    output, _, wb = books["messy"]
    ws = wb["Data_Quality"]
    values = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, 10)}
    assert values["Duplicate rows removed"] == 25
    all_text = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert "invalid_value_removed" in all_text and "duplicate_removed" in all_text


def test_ai_sheet_only_with_ai_output(tmp_path):
    raw, report = load_raw(SAMPLE_DIR / "test_b_medium.csv")
    output = run_pipeline(prepare(raw, report), db_path=tmp_path / "a.db")
    ai = {"recommendations": [{"text": "Check the landing page.", "evidence_ids": ["E01"],
                               "priority": "high", "metric_to_watch": "CPL"}]}
    sheets = generate_workbook(output.analysis, tmp_path / "ai.xlsx", ai=ai)
    assert "AI_Insights" in sheets
    ws = load_workbook(tmp_path / "ai.xlsx")["AI_Insights"]
    assert "AI-GENERATED" in ws["A1"].value


def test_export_saves_and_records(tmp_path):
    db = tmp_path / "e.db"
    raw, report = load_raw(SAMPLE_DIR / "test_c_large.csv")
    output = run_pipeline(prepare(raw, report), db_path=db)
    path = export_workbook(output.analysis, output.clean.clean_df, output.clean.quality_log_df,
                           exports_dir=tmp_path / "exports", db_path=db)
    assert path.exists() and path.parent.name == str(output.run_id) and path.suffix == ".xlsx"
    assert repo.list_reports(output.run_id, db_path=db)["report_type"].tolist() == ["excel"]


def test_reports_page_generates_excel(tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    monkeypatch.setattr(dbc, "settings", Settings(database_path=str(tmp_path / "app.db")))
    monkeypatch.setattr(excel_report, "EXPORTS_DIR", tmp_path / "exports")
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=240).run()
    at.button(key="load_sample").click().run()
    at.button(key="run_analysis").click().run()
    at.sidebar.radio(key="page").set_value("Reports").run()
    at.button(key="generate_excel").click().run()
    assert not at.exception and not at.error
    assert any("Workbook ready" in s.value for s in at.success)
    assert list((tmp_path / "exports").rglob("*.xlsx"))
