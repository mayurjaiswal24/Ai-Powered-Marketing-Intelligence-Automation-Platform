"""Interface polish: India-time display, creator signature and the grouped sidebar menu."""

from __future__ import annotations

import datetime as dt

import config.settings as cfg
from dashboard import layout, theme
from utils.formatting import format_datetime_ist


def test_ist_times_are_shown_in_india_time():
    assert format_datetime_ist("2026-09-27T12:04:00+00:00") == "27 Sep 2026, 5:34 PM"
    assert format_datetime_ist("2026-09-27T18:40:00") == "28 Sep 2026, 12:10 AM"   # no zone = UTC
    assert format_datetime_ist(dt.datetime(2026, 1, 5, 0, 0, tzinfo=dt.timezone.utc)) == "5 Jan 2026, 5:30 AM"
    assert format_datetime_ist(None) == "N/A" and format_datetime_ist("not a date") == "N/A"


def test_creator_details_and_signature():
    assert cfg.CREATOR_NAME == "Mayur Jaiswal"
    assert cfg.CREATOR_LINKEDIN.startswith("https://www.linkedin.com/in/")
    assert cfg.APP_VERSION == "1.0"
    assert cfg.REPORT_SIGNATURE == ("Prepared with Marketing Intelligence Platform · "
                                    "Created by Mayur Jaiswal")


def test_menu_groups_cover_every_page_once():
    assert layout.PAGES[0] == "Upload & Profile" and layout.PAGES[-1] == "About"
    assert len(layout.PAGES) == len(set(layout.PAGES)) == 12
    assert set(layout.PAGE_ICONS) == set(layout.PAGES) == set(layout.PAGE_DESCRIPTIONS)
    starts = layout._nav_section_starts()
    assert starts == {1: "Data", 3: "Analysis", 10: "AI", 11: "Output", 12: ""}
    css = theme.nav_css(starts)
    assert "nth-child(3)" in css and "'ANALYSIS'" in css


def _sample_output(tmp_path, monkeypatch):
    import database.connection as dbc
    from config.settings import Settings
    from dashboard.pipeline import DEFAULT_SAMPLE, load_raw, prepare, run_pipeline, sample_path
    monkeypatch.setattr(dbc, "settings", Settings(database_path=str(tmp_path / "app.db")))
    raw, report = load_raw(sample_path(DEFAULT_SAMPLE))
    return run_pipeline(prepare(raw, report))


def test_pdf_and_excel_carry_the_signature(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from pypdf import PdfReader
    from reports.excel_report import generate_workbook
    from reports.pdf_report import generate_pdf
    output = _sample_output(tmp_path, monkeypatch)
    xlsx, pdf = tmp_path / "book.xlsx", tmp_path / "report.pdf"
    generate_workbook(output.analysis, xlsx)
    first = load_workbook(xlsx).worksheets[0]
    assert cfg.REPORT_SIGNATURE in [first.cell(r, 1).value for r in range(1, 4)]
    assert str(first.cell(8, 2).value).endswith(" IST")            # "Generated" row in India time
    generate_pdf(output.analysis, pdf)
    pages = [p.extract_text() or "" for p in PdfReader(pdf).pages]
    assert all("Created by Mayur Jaiswal" in text for text in pages)   # cover + every footer
