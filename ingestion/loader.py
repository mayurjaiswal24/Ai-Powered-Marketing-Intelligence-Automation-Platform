"""File loading: turn an uploaded CSV/XLSX into a raw table of text, or fail politely.

Everything is loaded as TEXT on purpose. The original values (e.g. "₹1,23,456", "N/A",
"05/01/2026") must survive untouched so the data-quality log in Phase 04 can show the user
exactly what was in their file. Converting to numbers and dates happens later, in cleaning.

Alternative considered: `pandas.read_csv` / `read_excel` directly. We parse rows ourselves
(csv module / openpyxl) so CSV and Excel share one header-detection routine and so Excel dates
come out as clean text instead of "2026-05-10 00:00:00".
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from config.settings import settings

SUPPORTED_EXTENSIONS = {".csv": "csv", ".xlsx": "xlsx"}
CSV_ENCODINGS = ["utf-8", "utf-8-sig", "latin-1"]
CSV_DELIMITERS = {",": "comma", ";": "semicolon", "\t": "tab"}
HEADER_SEARCH_ROWS = 30  # how far down we look for the real header row

# Every .xlsx is a ZIP archive; an encrypted (password-protected) workbook is instead an
# old-style "OLE" container holding a stream called EncryptedPackage.
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ENCRYPTED_MARKER = "EncryptedPackage".encode("utf-16-le")


class IngestionError(Exception):
    """A file could not be loaded. `user_message` is safe to show in the UI as-is."""

    def __init__(self, user_message: str, code: str = "unreadable"):
        super().__init__(user_message)
        self.user_message = user_message
        self.code = code


@dataclass
class LoadReport:
    filename: str
    file_type: str               # "csv" or "xlsx"
    file_hash: str               # SHA-256 of the raw bytes (detects re-uploads)
    file_size_bytes: int
    rows: int = 0
    columns: int = 0
    sheet_used: str | None = None
    other_sheets: list[str] = field(default_factory=list)
    encoding: str | None = None
    delimiter: str | None = None  # "comma", "semicolon" or "tab"
    header_row: int = 1           # 1-based row number of the header in the original file
    warnings: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------------------------

def load_file(source: Any, filename: str | None = None, *, max_file_mb: int | None = None,
              large_row_warning: int | None = None) -> tuple[pd.DataFrame, LoadReport]:
    """Load a CSV/XLSX from a file path or a Streamlit UploadedFile (anything with
    `.getvalue()` or `.read()`). Returns (raw_df with every cell as text, LoadReport).
    Raises IngestionError with a plain-English message if the file cannot be used."""
    max_file_mb = max_file_mb or settings.max_upload_mb
    large_row_warning = large_row_warning or settings.large_row_warning

    filename = filename or _guess_name(source)
    extension = Path(filename).suffix.lower()
    if extension == ".xls":
        raise IngestionError(
            "This is an old Excel format (.xls). Please open it in Excel, save it as "
            "'Excel Workbook (.xlsx)', and upload it again.", "unsupported_type")
    if extension not in SUPPORTED_EXTENSIONS:
        shown = extension or "files without an extension"
        raise IngestionError(
            f"{shown} is not supported. Please upload a CSV (.csv) or Excel (.xlsx) file.",
            "unsupported_type")
    file_type = SUPPORTED_EXTENSIONS[extension]

    data = _read_bytes(source, max_file_mb)
    if not data.strip():
        raise IngestionError("The file is empty. Please check that you selected the right file.",
                             "empty")

    report = LoadReport(filename=filename, file_type=file_type,
                        file_hash=hashlib.sha256(data).hexdigest(), file_size_bytes=len(data))

    if file_type == "csv":
        rows = _csv_rows(data, report)
    else:
        rows = _xlsx_rows(data, report)

    df = _build_frame(rows, report)
    report.rows, report.columns = df.shape
    if report.rows > large_row_warning:
        report.warnings.append(
            f"This file has {report.rows:,} rows. Analysis may take a little longer than usual.")
    return df, report


# ---------------------------------------------------------------------------------------------
# Reading bytes
# ---------------------------------------------------------------------------------------------

def _guess_name(source: Any) -> str:
    if isinstance(source, (str, Path)):
        return Path(source).name
    name = getattr(source, "name", None)
    if name:
        return Path(str(name)).name
    raise IngestionError("The file name is missing, so the file type cannot be identified.",
                         "unsupported_type")


def _read_bytes(source: Any, max_file_mb: int) -> bytes:
    limit = max_file_mb * 1024 * 1024
    too_big = (f"The file is larger than the {max_file_mb} MB limit. Please upload a smaller "
               "file, or split it into parts.")

    if isinstance(source, (str, Path)):
        path = Path(source)
        if not path.is_file():
            raise IngestionError("The file could not be found. Please choose it again.", "not_found")
        if path.stat().st_size > limit:
            raise IngestionError(too_big, "too_large")
        return path.read_bytes()

    size = getattr(source, "size", None)  # Streamlit UploadedFile knows its size up front
    if isinstance(size, int) and size > limit:
        raise IngestionError(too_big, "too_large")
    if hasattr(source, "getvalue"):
        data = source.getvalue()
    elif hasattr(source, "read"):
        data = source.read()
    else:
        raise IngestionError("This upload could not be read. Please try uploading it again.")
    if isinstance(data, str):
        data = data.encode("utf-8")
    if len(data) > limit:
        raise IngestionError(too_big, "too_large")
    return bytes(data)


# ---------------------------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------------------------

def _decode(data: bytes, report: LoadReport) -> str:
    if b"\x00" in data[:4096]:
        # Text files never contain NUL bytes; this is binary data with a .csv name.
        raise IngestionError("This file could not be read. It may be damaged or not a real "
                             "CSV text file.", "corrupt")
    if data.startswith(b"\xef\xbb\xbf"):
        report.encoding = "utf-8-sig"  # UTF-8 with the byte-order mark Excel adds
        return data.decode("utf-8-sig")
    for encoding in CSV_ENCODINGS:
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        report.encoding = encoding
        if encoding == "latin-1":
            report.warnings.append(
                "The file is not UTF-8, so it was read as Latin-1. A few special characters "
                "(such as ₹) may not display correctly.")
        return text
    raise IngestionError("This file could not be read as text.", "corrupt")  # pragma: no cover


def _detect_delimiter(text: str) -> str:
    """Pick the delimiter that splits the first lines into the most, and most consistent,
    columns. The csv module respects quotes, so "1,23,456" inside quotes is not a split."""
    sample = "\n".join(text.splitlines()[:50])
    best, best_score = ",", (-1.0, 0)
    for delimiter in CSV_DELIMITERS:
        lines = [r for r in csv.reader(io.StringIO(sample), delimiter=delimiter) if any(c.strip() for c in r)]
        if not lines:
            continue
        widths = [len(r) for r in lines]
        typical = max(set(widths), key=widths.count)
        if typical < 2:
            continue
        consistency = widths.count(typical) / len(widths)
        score = (consistency, typical)
        if score > best_score:
            best, best_score = delimiter, score
    return best


def _csv_rows(data: bytes, report: LoadReport) -> list[list[str]]:
    text = _decode(data, report)
    delimiter = _detect_delimiter(text)
    report.delimiter = CSV_DELIMITERS[delimiter]
    try:
        return [row for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)]
    except csv.Error as exc:
        raise IngestionError("This CSV file could not be read because its structure is broken "
                             "(for example, a quote mark that is never closed).", "corrupt") from exc


# ---------------------------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------------------------

def _cell_text(value: Any) -> str:
    """Convert one Excel cell to text without losing what the user typed."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, dt.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _xlsx_rows(data: bytes, report: LoadReport) -> list[list[str]]:
    if data.startswith(OLE_MAGIC):
        if ENCRYPTED_MARKER in data:
            raise IngestionError(
                "This Excel file is password-protected. Please remove the password in Excel "
                "(File > Info > Protect Workbook > Encrypt with Password, then clear it) and "
                "upload it again.", "password_protected")
        raise IngestionError(
            "This looks like an old Excel format (.xls) renamed to .xlsx. Please open it in "
            "Excel, save it as 'Excel Workbook (.xlsx)', and upload it again.", "corrupt")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise IngestionError("This file could not be read. It may be damaged or not a real "
                             "Excel file.", "corrupt") from exc

    try:
        sheets: dict[str, list[list[str]]] = {}
        filled: dict[str, int] = {}
        for ws in workbook.worksheets:
            rows = [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
            sheets[ws.title] = rows
            filled[ws.title] = sum(1 for row in rows for c in row if c.strip())
    except Exception as exc:  # any failure while reading cells means a damaged workbook
        raise IngestionError("This file could not be read. It may be damaged or not a real "
                             "Excel file.", "corrupt") from exc
    finally:
        workbook.close()

    if not sheets or max(filled.values()) == 0:
        raise IngestionError("The Excel file has no data in any sheet.", "empty")

    # Choose the sheet with the most filled-in cells (ties -> the first sheet).
    chosen = max(sheets, key=lambda name: filled[name])
    report.sheet_used = chosen
    report.other_sheets = [name for name in sheets if name != chosen]
    if report.other_sheets:
        report.warnings.append(
            f"The workbook has {len(sheets)} sheets. Using '{chosen}' because it has the most "
            f"data. Other sheets not used: {', '.join(report.other_sheets)}.")
    return sheets[chosen]


# ---------------------------------------------------------------------------------------------
# Shared: header detection and DataFrame building
# ---------------------------------------------------------------------------------------------

def _filled(row: list[str]) -> int:
    return sum(1 for c in row if c.strip())


def _find_header(rows: list[list[str]]) -> int:
    """Index of the header row: the first row that is about as wide as the table itself.
    A title ("Campaign report, Q3") or blank lines above the table are therefore skipped."""
    top = rows[:HEADER_SEARCH_ROWS]
    widest = max((_filled(r) for r in top), default=0)
    needed = max(2, int(widest * 0.6 + 0.5)) if widest >= 2 else 1
    for i, row in enumerate(top):
        if _filled(row) >= needed:
            return i
    return 0


def _unique_headers(raw: list[str], report: LoadReport) -> list[str]:
    headers, seen = [], {}
    blank_count = 0
    for i, name in enumerate(raw, start=1):
        name = name if name.strip() else ""
        if not name:
            blank_count += 1
            name = f"Unnamed column {i}"
        if name in seen:
            seen[name] += 1
            new_name = f"{name} ({seen[name]})"
            report.warnings.append(f"Column name '{name}' appears more than once; the repeat "
                                   f"was renamed to '{new_name}'.")
            name = new_name
        else:
            seen[name] = 1
        headers.append(name)
    if blank_count:
        report.warnings.append(f"{blank_count} column(s) had no heading and were given a "
                               "placeholder name such as 'Unnamed column 3'.")
    return headers


def _build_frame(rows: list[list[str]], report: LoadReport) -> pd.DataFrame:
    if not any(_filled(r) for r in rows):
        raise IngestionError("The file is empty. Please check that you selected the right file.",
                             "empty")

    header_idx = _find_header(rows)
    report.header_row = header_idx + 1
    if header_idx > 0:
        report.warnings.append(
            f"Skipped {header_idx} row(s) above the column headings (a title or blank rows). "
            f"Headings were found on row {header_idx + 1}.")

    header = rows[header_idx]
    below = rows[header_idx + 1:]
    body = [r for r in below if _filled(r)]
    # Blank rows between data rows are worth mentioning; trailing blank lines are not.
    last_filled = max((i for i, r in enumerate(below) if _filled(r)), default=-1)
    interior_blank_rows = sum(1 for r in below[:last_filled] if not _filled(r))

    # Width = header width, extended if data rows are wider; trailing empty columns dropped.
    width = max([len(header)] + [len(r) for r in body])
    while width > 0 and _column_empty(header, body, width - 1):
        width -= 1
    if width == 0:
        raise IngestionError("The file is empty. Please check that you selected the right file.",
                             "empty")

    if not body:
        raise IngestionError(
            "The file has column headings but no data rows. Please upload a file that "
            "contains data under the headings.", "header_only")

    ragged = sum(1 for r in body if len(r) > len(header) and any(c.strip() for c in r[len(header):]))
    if ragged:
        report.warnings.append(f"{ragged} row(s) had more values than there are column headings; "
                               "the extra values were kept in unnamed columns.")
    if interior_blank_rows:
        report.warnings.append(f"Ignored {interior_blank_rows} completely empty row(s).")

    padded_header = (header + [""] * width)[:width]
    columns = _unique_headers(padded_header, report)
    data = [(r + [""] * width)[:width] for r in body]
    return pd.DataFrame(data, columns=columns, dtype=str)


def _column_empty(header: list[str], body: list[list[str]], col: int) -> bool:
    if col < len(header) and header[col].strip():
        return False
    return all(col >= len(r) or not r[col].strip() for r in body)
